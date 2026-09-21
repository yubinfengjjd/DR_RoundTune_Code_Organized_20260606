
from pathlib import Path
import textwrap, os, zipfile
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle, Ellipse, FancyArrowPatch, Rectangle
import matplotlib as mpl

ROOT = Path("chatgpt_exports")
ROOT.mkdir(exist_ok=True)
FIGDIR = ROOT / "ferroptosis_organelle_figures"
FIGDIR.mkdir(exist_ok=True)
OUT = ROOT / "Ferroptosis_Organelle_Crosstalk_Review_Framework_Figures_Tables_CN.docx"

mpl.rcParams["font.family"] = "DejaVu Sans"
mpl.rcParams["figure.dpi"] = 160

COLORS = {
    "blue":"#DCEBFA","blue2":"#82B5E8","green":"#DDF3E5","green2":"#58A87B",
    "red":"#F9DFE0","red2":"#D85E65","purple":"#ECE3FA","purple2":"#8A6DC1",
    "yellow":"#FFF0C9","yellow2":"#C58B27","teal":"#DFF4F1","teal2":"#4C9F95",
    "grey":"#EEF2F5","dark":"#17365D"
}

def clean_ax(ax, title):
    ax.set_xlim(0, 16); ax.set_ylim(0, 10); ax.axis("off")
    ax.text(0.2, 9.6, title, fontsize=18, fontweight="bold", color="#102B55", va="top")

def box(ax, xy, wh, title, bullets=None, fc="#F7F9FB", ec="#B6C6D6", tc="#17365D", fs=11):
    x,y=xy; w,h=wh
    p=FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.18,rounding_size=0.18",
                     linewidth=1.4,edgecolor=ec,facecolor=fc)
    ax.add_patch(p)
    ax.text(x+0.18,y+h-0.25,title,fontsize=fs+1,fontweight="bold",color=tc,va="top")
    if bullets:
        yy=y+h-0.75
        for b in bullets:
            lines=textwrap.wrap(b, width=max(18,int(w*8)))
            ax.text(x+0.28,yy,u"\u2022 "+lines[0],fontsize=fs-1,color="#233445",va="top")
            yy-=0.36
            for ln in lines[1:]:
                ax.text(x+0.52,yy,ln,fontsize=fs-1,color="#233445",va="top")
                yy-=0.34
            yy-=0.05
    return p

def arrow(ax, a, b, color="#5A7896", lw=2.2, rad=0):
    ax.add_patch(FancyArrowPatch(a,b,arrowstyle="-|>",mutation_scale=18,
                                 linewidth=lw,color=color,
                                 connectionstyle=f"arc3,rad={rad}"))

def organelle(ax, center, label, fc, ec, kind="ellipse"):
    x,y=center
    if kind=="ellipse":
        p=Ellipse((x,y),2.2,1.25,facecolor=fc,edgecolor=ec,linewidth=2)
    elif kind=="circle":
        p=Circle((x,y),0.8,facecolor=fc,edgecolor=ec,linewidth=2)
    else:
        p=FancyBboxPatch((x-1.1,y-0.6),2.2,1.2,boxstyle="round,pad=0.15,rounding_size=0.25",facecolor=fc,edgecolor=ec,linewidth=2)
    ax.add_patch(p)
    ax.text(x,y,label,ha="center",va="center",fontsize=11,fontweight="bold",color="#24364B")
    return p

def save_fig(fig, name):
    path=FIGDIR/name
    fig.savefig(path,bbox_inches="tight",facecolor="white")
    plt.close(fig)
    return path

# Figure 1
fig,ax=plt.subplots(figsize=(14,9))
clean_ax(ax,"Figure 1 | Organelle crosstalk as the organizing principle of ferroptosis")
ax.add_patch(Circle((8,5.2),1.25,facecolor="#FCE7E3",edgecolor="#D85E65",linewidth=2.5))
ax.text(8,5.35,"FERROPTOSIS",ha="center",va="center",fontsize=17,fontweight="bold",color="#9E2F35")
ax.text(8,4.85,"iron + PUFA-PL peroxidation\\n+ failed antioxidant defense",ha="center",va="center",fontsize=9)
nodes=[
((2.0,7.4),"Mitochondria",COLORS["red"],COLORS["red2"],["TCA/ETC & ROS","Fe-S / iron","DHODH-CoQH2"]),
((7.8,8.1),"Endoplasmic reticulum",COLORS["blue"],COLORS["blue2"],["ACSL4/LPCAT3","UPR & Ca2+","MAM contacts"]),
((13.3,7.2),"Lysosome / autophagy",COLORS["purple"],COLORS["purple2"],["NCOA4 ferritinophagy","lipophagy","lysosomal redox"]),
((13.2,3.0),"Peroxisome",COLORS["green"],COLORS["green2"],["PUFA ether lipids","VLCFA handling","ROS metabolism"]),
((8.0,1.5),"Golgi / plasma membrane",COLORS["blue"],COLORS["blue2"],["lipid trafficking","FSP1-CoQ","terminal membrane failure"]),
((2.1,2.2),"Lipid droplets",COLORS["yellow"],COLORS["yellow2"],["PUFA storage","lipolysis","buffer vs fuel"]),
((3.9,5.0),"Nucleus",COLORS["purple"],COLORS["purple2"],["NRF2, p53, ATF4","stress adaptation","cell-state control"]),
]
for (c,l,fc,ec,buls) in nodes:
    organelle(ax,c,l,fc,ec)
    # mini bullet box close by
for p in [(3.1,7.1),(7.8,7.0),(12.0,6.2),(12.1,2.1),(7.8,2.7),(3.0,3.2),(4.9,4.0)]:
    pass
for c,l,fc,ec,buls in nodes:
    arrow(ax,c,(8,5.2),ec,1.8)
# protective systems
box(ax,(0.4,8.2),(3.5,1.1),"Protective brakes",["GPX4-GSH | FSP1-CoQ | DHODH | GCH1-BH4"],fc="#E8F6EC",ec="#7DBD8D",fs=10)
box(ax,(5.4,0.15),(5.2,0.9),"Integrated outcome",["Homeostasis -> priming/sensitization -> membrane failure/death"],fc="#F3F5F7",ec="#9AA9B8",fs=10)
f1=save_fig(fig,"figure1_overview.png")

# Figure 2
fig,ax=plt.subplots(figsize=(14,9)); clean_ax(ax,"Figure 2 | Mitochondria-ER-peroxisome metabolic circuits that tune ferroptosis")
organelle(ax,(3.1,5.5),"Mitochondria",COLORS["red"],COLORS["red2"])
organelle(ax,(8.0,7.1),"ER",COLORS["blue"],COLORS["blue2"])
organelle(ax,(12.9,5.3),"Peroxisome",COLORS["green"],COLORS["green2"])
for a,b,col in [((4.2,5.9),(6.8,6.8),COLORS["red2"]),((9.2,6.7),(11.8,5.7),COLORS["green2"]),((11.9,4.8),(4.2,5.1),"#6C7FB3")]:
    arrow(ax,a,b,col,2.4)
box(ax,(0.5,1.1),(4.5,2.3),"Mitochondrial module",["TCA / ETC metabolic flux","mitochondrial ROS and iron","Fe-S cluster homeostasis","DHODH-CoQH2 membrane defense"],fc=COLORS["red"],ec=COLORS["red2"])
box(ax,(5.7,1.1),(4.5,2.3),"ER module",["ACSL4/LPCAT3 PUFA-PL remodeling","ER stress / UPR","Ca2+ storage and transfer","MAMs coordinate lipid/redox flux"],fc=COLORS["blue"],ec=COLORS["blue2"])
box(ax,(10.9,1.1),(4.5,2.3),"Peroxisomal module",["PUFA ether-phospholipid synthesis","VLCFA metabolism","H2O2 / redox metabolism","ER cooperation for membrane lipid supply"],fc=COLORS["green"],ec=COLORS["green2"])
ax.add_patch(Circle((8,4.5),0.95,facecolor="#FCE7E3",edgecolor="#D85E65",linewidth=2))
ax.text(8,4.5,"ferroptotic\\nthreshold",ha="center",va="center",fontsize=13,fontweight="bold",color="#9E2F35")
arrow(ax,(3.9,5.0),(7.0,4.7),COLORS["red2"],2); arrow(ax,(8,6.0),(8,5.45),COLORS["blue2"],2); arrow(ax,(12.0,5.0),(9.0,4.7),COLORS["green2"],2)
f2=save_fig(fig,"figure2_metabolic_circuits.png")

# Figure 3
fig,ax=plt.subplots(figsize=(14,9)); clean_ax(ax,"Figure 3 | Lysosome-autophagy-lipid droplet crosstalk mobilizes iron and lipids")
organelle(ax,(7.9,5.1),"Lysosome",COLORS["purple"],COLORS["purple2"],kind="circle")
organelle(ax,(3.0,6.4),"Ferritin",COLORS["blue"],COLORS["blue2"],kind="circle")
organelle(ax,(3.0,3.2),"Autophagosome",COLORS["blue"],COLORS["blue2"],kind="ellipse")
organelle(ax,(12.7,5.2),"Lipid droplets",COLORS["yellow"],COLORS["yellow2"],kind="ellipse")
organelle(ax,(8.1,2.0),"Mitochondria",COLORS["red"],COLORS["red2"],kind="ellipse")
arrow(ax,(3.8,6.2),(7.0,5.5),COLORS["purple2"],2.3); ax.text(4.8,6.45,"NCOA4 ferritinophagy",fontsize=10,color="#6C4A93")
arrow(ax,(3.9,3.5),(7.0,4.7),COLORS["blue2"],2.3); ax.text(4.5,3.65,"selective autophagy",fontsize=10,color="#346CA8")
arrow(ax,(11.8,5.4),(8.9,5.3),COLORS["yellow2"],2.3); ax.text(10.1,5.8,"lipophagy",fontsize=10,color="#8B6416")
arrow(ax,(8.0,3.0),(8.0,4.2),COLORS["red2"],2.1); ax.text(8.3,3.4,"ROS / damage",fontsize=10,color="#A44246")
box(ax,(0.45,7.7),(4.1,1.2),"Iron mobilization",["Ferritin -> lysosome -> Fe2+ -> labile iron pool"],fc=COLORS["red"],ec=COLORS["red2"],fs=10)
box(ax,(5.3,7.7),(4.8,1.2),"Antioxidant turnover",["Chaperone-mediated autophagy can alter GPX4/protein stability"],fc=COLORS["purple"],ec=COLORS["purple2"],fs=10)
box(ax,(10.8,7.7),(4.7,1.2),"Lipid mobilization",["Storage protects; lipolysis/lipophagy releases PUFA substrate"],fc=COLORS["yellow"],ec=COLORS["yellow2"],fs=10)
box(ax,(5.2,0.2),(5.7,1.0),"Convergence",["Fe2+ + PUFA availability + redox stress -> lipid peroxidation -> ferroptosis"],fc="#FCE7E3",ec=COLORS["red2"],fs=10)
f3=save_fig(fig,"figure3_lysosome_autophagy_ld.png")

# Figure 4
fig,ax=plt.subplots(figsize=(14,9)); clean_ax(ax,"Figure 4 | Nucleus, Golgi and plasma membrane coordinate adaptation and execution")
organelle(ax,(3.0,6.4),"Nucleus",COLORS["purple"],COLORS["purple2"],kind="circle")
organelle(ax,(8.1,6.4),"Golgi",COLORS["blue"],COLORS["blue2"],kind="ellipse")
ax.add_patch(Rectangle((2.1,2.0),11.8,0.55,facecolor="#D6EFE2",edgecolor="#4A9A74",linewidth=2))
ax.text(8.0,2.72,"Plasma membrane",ha="center",fontsize=12,fontweight="bold",color="#2D6D52")
box(ax,(0.6,4.0),(4.7,1.5),"Transcriptional adaptation",["NRF2 | p53 | ATF4 | HIFs | YAP/TAZ","SLC7A11 / GPX4 / ferritin / ACSL4 programs"],fc=COLORS["purple"],ec=COLORS["purple2"],fs=10)
box(ax,(5.8,4.0),(4.6,1.5),"Trafficking and lipid supply",["Golgi stress","delivery of membrane lipids and protective proteins"],fc=COLORS["blue"],ec=COLORS["blue2"],fs=10)
box(ax,(10.9,4.0),(4.5,1.5),"Terminal membrane defense",["system Xc-","FSP1-CoQ","membrane repair / ion homeostasis"],fc=COLORS["green"],ec=COLORS["green2"],fs=10)
arrow(ax,(3.7,5.9),(6.8,5.9),COLORS["purple2"],2); arrow(ax,(8.2,5.7),(8.2,2.65),COLORS["blue2"],2)
ax.add_patch(Circle((8.0,2.25),0.45,facecolor="#FA8A83",edgecolor="#B23535",linewidth=2))
ax.text(8.0,2.25,"LPO",ha="center",va="center",fontsize=10,fontweight="bold",color="white")
ax.text(8.0,1.25,"adaptation -> threshold -> catastrophic membrane failure",ha="center",fontsize=13,fontweight="bold",color="#9E2F35")
f4=save_fig(fig,"figure4_execution.png")

# Figure 5
fig,ax=plt.subplots(figsize=(14,9)); clean_ax(ax,"Figure 5 | Organelle-defined ferroptosis endotypes across disease settings")
diseases=[
("Cancer","ER + mitochondria + lipid droplets","PUFA remodeling / antioxidant escape",COLORS["red"],COLORS["red2"]),
("Ischaemia-reperfusion","Mitochondria + membrane","ROS burst / lipid oxidation",COLORS["blue"],COLORS["blue2"]),
("Acute kidney injury","Mitochondria + lysosome","iron mobilization / tubular injury",COLORS["green"],COLORS["green2"]),
("Neurodegeneration","Mitochondria + ER + lysosome","neuronal redox / iron stress",COLORS["purple"],COLORS["purple2"]),
("Liver disease","ER + lipid droplets + peroxisome","lipid overload / PUFA supply",COLORS["yellow"],COLORS["yellow2"]),
("Intestinal injury / IBD","Plasma membrane + mitochondria + ER","barrier lipid peroxidation",COLORS["teal"],COLORS["teal2"]),
]
coords=[(0.5,5.5),(5.35,5.5),(10.2,5.5),(0.5,2.1),(5.35,2.1),(10.2,2.1)]
for (name,axis,mech,fc,ec),(x,y) in zip(diseases,coords):
    box(ax,(x,y),(4.3,2.5),name,[axis,mech,"Therapeutic direction is context-dependent"],fc=fc,ec=ec,fs=10)
ax.text(8,1.0,"Disease classification by dominant organelle dependency rather than by ferroptosis label alone",
        ha="center",fontsize=13,fontweight="bold",color="#17365D")
f5=save_fig(fig,"figure5_disease_endotypes.png")

# Figure 6
fig,ax=plt.subplots(figsize=(14,9)); clean_ax(ax,"Figure 6 | Therapeutic targeting of organelle crosstalk in ferroptosis")
layers=[
("1. Iron handling",["Chelators","TFRC / ferritin","NCOA4 ferritinophagy"],COLORS["blue"],COLORS["blue2"]),
("2. Antioxidant systems",["GPX4-GSH","FSP1-CoQ","DHODH-CoQH2","GCH1-BH4"],COLORS["green"],COLORS["green2"]),
("3. Lipid remodeling",["ACSL4 / LPCAT3","MUFA-PUFA balance","ether-lipid metabolism"],COLORS["yellow"],COLORS["yellow2"]),
("4. Organelle stress",["ER-UPR","mitochondrial quality control","lysosome/autophagy","peroxisomal metabolism"],COLORS["purple"],COLORS["purple2"]),
("5. Membrane protection",["radical-trapping antioxidants","ferrostatin-like agents","membrane repair"],COLORS["red"],COLORS["red2"]),
]
x=0.35
for title,buls,fc,ec in layers:
    box(ax,(x,4.2),(2.8,4.1),title,buls,fc=fc,ec=ec,fs=9)
    x+=3.1
ax.add_patch(FancyBboxPatch((0.8,1.25),6.7,1.8,boxstyle="round,pad=0.18,rounding_size=0.2",facecolor="#F7E7E8",edgecolor="#D85E65",linewidth=1.7))
ax.text(4.15,2.72,"Cancer: induce ferroptosis selectively",ha="center",fontsize=12,fontweight="bold",color="#9E2F35")
ax.text(4.15,2.15,"exploit tumour-specific organelle dependencies\\nand antioxidant escape mechanisms",ha="center",fontsize=10)
ax.add_patch(FancyBboxPatch((8.5,1.25),6.7,1.8,boxstyle="round,pad=0.18,rounding_size=0.2",facecolor="#E7F1FA",edgecolor="#82B5E8",linewidth=1.7))
ax.text(11.85,2.72,"Tissue injury / degeneration: inhibit ferroptosis",ha="center",fontsize=12,fontweight="bold",color="#2D629B")
ax.text(11.85,2.15,"preserve organelle quality, iron balance\\nand membrane integrity",ha="center",fontsize=10)
ax.text(8,0.55,"Biomarkers + cell-state specificity + organelle-targeted delivery + combination therapy",
        ha="center",fontsize=12,fontweight="bold",color="#17365D")
f6=save_fig(fig,"figure6_therapy.png")

# ---------- DOCX ----------
doc = Document()
sec = doc.sections[0]
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)

# styles
for name in ["Normal","Title","Heading 1","Heading 2","Heading 3"]:
    st=doc.styles[name]
    st.font.name="Times New Roman"
    st._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
doc.styles["Normal"].font.size=Pt(10.5)
doc.styles["Title"].font.size=Pt(20)
doc.styles["Heading 1"].font.size=Pt(15)
doc.styles["Heading 2"].font.size=Pt(12.5)
doc.styles["Heading 3"].font.size=Pt(11)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn("w:shd"))
    if shd is None:
        shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)

def run_font(run, size=10.5, bold=False, color=None, east="宋体"):
    run.font.name="Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"),east)
    run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=RGBColor.from_string(color)

def para(text="", bold=False, center=False, size=10.5, indent=True, space=5):
    p=doc.add_paragraph()
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if indent: p.paragraph_format.first_line_indent=Cm(0.74)
    p.paragraph_format.space_after=Pt(space)
    r=p.add_run(text); run_font(r,size,bold)
    return p

def bullet(text):
    p=doc.add_paragraph()
    p.paragraph_format.left_indent=Cm(0.55)
    p.paragraph_format.first_line_indent=Cm(-0.3)
    p.paragraph_format.space_after=Pt(2)
    r=p.add_run("• "+text); run_font(r,10.2)
    return p

def heading(text, level=1):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10 if level==1 else 6)
    p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text)
    if level==1: run_font(r,15,True,"17365D","黑体")
    elif level==2: run_font(r,12.5,True,"245B78","黑体")
    else: run_font(r,11,True,"3B6A57","黑体")
    return p

def add_figure(path,title,legend):
    doc.add_picture(str(path),width=Inches(6.55))
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(title); run_font(r,9.8,True,"17365D")
    p2=doc.add_paragraph(); p2.paragraph_format.space_after=Pt(8)
    r2=p2.add_run(legend); run_font(r2,9.2)

def add_table(title, headers, rows):
    heading(title,2)
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c,"DCEAF5")
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(h); run_font(r,8.8,True,"17365D","黑体")
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(str(val)); run_font(r,8.4)
    doc.add_paragraph()

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("细胞器与铁死亡串扰"); run_font(r,20,True,"17365D","黑体")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("从亚细胞网络到疾病内表型与精准干预"); run_font(r,14,True,"245B78","黑体")
para("Organelle crosstalk in ferroptosis: from subcellular networks to disease endotypes and precision targeting",
     center=True,size=11,indent=False)
para("综述框架、完整插图与可编辑表格整合版｜定位：影响因子 >10，CNS子刊风格",center=True,size=10,indent=False)

doc.add_page_break()

heading("摘要",1)
para("铁死亡（ferroptosis）是一种由铁依赖性磷脂过氧化驱动的调控性细胞死亡形式。传统研究主要围绕游离铁积累、GPX4失活、system Xc⁻抑制以及ACSL4依赖性多不饱和脂肪酸磷脂重塑展开。然而，越来越多证据表明，铁死亡并非发生于单一亚细胞区室，也不能由某一条线性信号通路充分解释。线粒体、内质网、溶酶体、自噬体、过氧化物酶体、脂滴、细胞核、高尔基体以及质膜分别控制铁处理、脂质供给、氧化还原稳态、膜组成、选择性自噬、应激转录以及终末膜损伤，并通过膜接触位点、脂质转运、Ca²⁺交换、活性氧扩散以及铁动员构成高度耦合的亚细胞网络。")
para("本文建议以“细胞器串扰”而非“单分子通路”作为全文组织原则：首先讨论不同细胞器的功能模块，随后解析线粒体–ER–过氧化物酶体代谢回路、溶酶体–自噬–脂滴铁脂动员轴，以及细胞核–Golgi–质膜的应激适应与终末执行轴；进一步提出“细胞器定义的铁死亡内表型”，并将其用于癌症、缺血再灌注、急性肾损伤、神经退行性疾病、肝病和肠道炎症等疾病的机制分层。最后讨论铁处理、GPX4/FSP1/DHODH等平行抗氧化系统、脂质重塑、选择性自噬、细胞器应激和细胞器靶向递送的治疗潜力。")
para("关键词：ferroptosis；organelle crosstalk；mitochondria；endoplasmic reticulum；lysosome；peroxisome；lipid droplet；lipid peroxidation",indent=False)

heading("1. Introduction：铁死亡是一个空间化组织的细胞死亡过程",1)
para("铁死亡通常被定义为由铁依赖性膜磷脂过氧化驱动的调控性细胞死亡。其经典分子框架包括三个彼此交织的核心模块：铁稳态失衡、多不饱和脂肪酸磷脂的形成和氧化，以及抗氧化防御系统崩溃。这一框架解释了铁死亡“需要什么”，却不足以解释铁死亡“在哪里发生”“从哪里开始”“如何在细胞内部传播”以及“为什么不同细胞对同一铁死亡诱导剂表现出截然不同的敏感性”。")
para("细胞内铁、脂质和抗氧化能力都具有强烈的区室化特征。因此真正决定铁死亡的不只是总体ROS水平，而是哪些脂质在何种细胞器中被氧化，以及这种局部损伤能否跨过细胞器和膜系统传播到不可逆阶段。")
para("全文建议主线：organelle homeostasis → inter-organelle communication → iron/lipid/redox redistribution → compartment-specific lipid peroxidation → membrane damage propagation → ferroptotic execution → disease-specific organelle endotypes → organelle-targeted therapy。",bold=True)

add_figure(f1,"Figure 1 | Organelle crosstalk as the organizing principle of ferroptosis",
"Ferroptosis emerges from coordinated dysfunction across multiple organelles. Mitochondria regulate metabolic and redox flux, the ER supplies peroxidation-prone membrane lipids, lysosomes and autophagy mobilize iron and lipids, peroxisomes shape specialized lipid metabolism, lipid droplets buffer or release fatty acids, the nucleus controls adaptive transcription, and Golgi/plasma-membrane systems coordinate trafficking and terminal membrane integrity.")

heading("2. 从分子通路到细胞器中心化框架",1)
para("铁死亡至少需要三个彼此交叉的条件：可反应铁增加、可过氧化膜磷脂积累，以及膜脂过氧化修复能力不足。GPX4–GSH是核心保护轴，但FSP1–CoQ、DHODH–CoQH₂和GCH1–BH4等平行系统说明，不同亚细胞区室具有相对独立的抗铁死亡防线。")
para("因此，高水平综述应把问题从“哪个分子控制铁死亡”升级为“哪个亚细胞区室首先突破其局部防御阈值，以及这种损伤如何通过细胞器串扰传播”。")

heading("3. 线粒体：代谢放大器还是情境依赖性参与者？",1)
para("线粒体参与TCA循环、电子传递、Fe–S簇和血红素代谢，并可产生影响铁死亡易感性的ROS。然而，线粒体并非所有铁死亡模式中的绝对必需细胞器，更适合被定义为情境依赖性的代谢与氧化放大器。")
bullet("TCA/ETC与线粒体ROS可在特定营养和诱导条件下提高易感性。")
bullet("线粒体铁和Fe–S簇稳态连接铁代谢、电子传递和氧化压力。")
bullet("DHODH–CoQH₂构成线粒体内膜的局部抗铁死亡防线，与GPX4形成互补。")

heading("4. 内质网：易过氧化膜脂的重塑枢纽",1)
para("内质网是磷脂合成与膜脂重塑中心。ACSL4促进PUFA形成acyl-CoA，LPCAT3参与其进入膜磷脂，从而直接塑造可发生过氧化的膜底物库。与此同时，ER stress/UPR、Ca²⁺稳态和ER–mitochondria接触共同决定应激是被缓冲还是转向死亡。")
bullet("ACSL4/LPCAT3决定PUFA-PL底物可用性。")
bullet("PERK、IRE1和ATF6等UPR模块具有时间、剂量和细胞状态依赖性。")
bullet("MAMs整合脂质转运、Ca²⁺流和线粒体代谢，是重要的空间调控节点。")

heading("5. 过氧化物酶体：特殊脂质合成与铁死亡可塑性",1)
para("过氧化物酶体参与VLCFA代谢、醚脂生物合成和氧化还原反应。需要避免把plasmalogen或ether lipid统一描述为抗氧化保护物：PUFA-containing ether phospholipids在特定细胞背景下可成为高效的铁死亡过氧化底物，因此其作用取决于脂质物种、酰基链构成、膜定位与细胞状态。")

add_figure(f2,"Figure 2 | Mitochondria–ER–peroxisome metabolic circuits that tune ferroptosis",
"Mitochondria, ER and peroxisomes form a metabolically connected network controlling ferroptosis. Mitochondria influence respiratory and iron/redox states; the ER drives PUFA-phospholipid remodeling and stress/Ca²⁺ signaling; and peroxisomes contribute specialized ether-lipid synthesis and oxidative metabolism. Contact sites couple these processes and determine whether metabolic stress remains adaptive or crosses a ferroptotic threshold.")

heading("6. 溶酶体–自噬轴：铁池与氧化底物的动员中心",1)
para("溶酶体是连接铁储存、铁释放和氧化损伤的重要细胞器。NCOA4介导ferritinophagy，将ferritin送入溶酶体降解并扩大labile iron pool；同时，lipophagy和其他选择性自噬可改变脂质供给、线粒体质量和抗氧化蛋白稳定性。")
para("因此不应笼统写成“autophagy promotes ferroptosis”。更严谨的框架是：不同cargo-selective autophagy通过铁、脂质、线粒体和蛋白稳态重编程铁死亡易感性。")

heading("7. 脂滴：PUFA缓冲池还是铁死亡燃料库？",1)
para("脂滴具有典型双相作用。将PUFA储存在TAG中可减少其直接进入膜磷脂，从而发挥短期缓冲作用；但当lipolysis或lipophagy增强时，脂肪酸重新释放，并可进入β-oxidation、膜磷脂合成或脂质过氧化路径。因此脂滴状态决定其是保护性仓库还是促铁死亡燃料池。")

add_figure(f3,"Figure 3 | Lysosome–autophagy–lipid droplet crosstalk mobilizes iron and lipids for ferroptosis",
"Lysosomes, selective autophagy and lipid droplets jointly control the availability of redox-active iron and peroxidizable lipids. NCOA4-mediated ferritinophagy mobilizes ferritin iron, whereas lipophagy and lipolysis release fatty-acid substrates. Lipid droplets can therefore buffer PUFAs or, when mobilized, fuel ferroptotic membrane remodeling.")

heading("8. 细胞核：铁死亡能力的转录控制器",1)
para("细胞核并不直接执行铁死亡，却决定细胞在特定应激背景下是否具备进入铁死亡的分子能力。NRF2、p53、ATF4、HIFs、YAP/TAZ和MTF1等转录调节因子将氧化、营养、缺氧和机械信号整合进铁、脂质和抗氧化程序。")
bullet("NRF2通常增强抗氧化和铁稳态程序，但在肿瘤中也可能形成治疗耐受。")
bullet("p53具有高度情境依赖的促/抗铁死亡作用，不应简单归纳为单向调节。")
bullet("ATF4、HIFs和YAP/TAZ说明铁死亡敏感性是动态的cell-state property。")

heading("9. Golgi与质膜：从物流适应到终末执行",1)
para("Golgi位于膜脂和蛋白质运输中心，可能通过膜组分更新、转运蛋白定位和应激反应影响铁死亡，但目前更适合被视为ferroptosis-supporting trafficking hub，而非经典执行器。")
para("质膜则是铁死亡终末阶段最关键的膜系统之一。FSP1–CoQ提供与GPX4并行的膜脂自由基抑制机制；当脂质过氧化传播超过膜修复和抗氧化能力后，膜通透性与Na⁺/Ca²⁺稳态崩溃，铁死亡进入不可逆阶段。")

add_figure(f4,"Figure 4 | Nucleus, Golgi and plasma membrane coordinate stress adaptation and terminal execution",
"Nuclear transcriptional programs control ferroptotic competence, the Golgi regulates lipid/protein trafficking, and the plasma membrane contains system Xc⁻ and FSP1–CoQ-dependent protective machinery. When adaptive programs are overwhelmed, lipid-peroxide propagation and ion imbalance drive catastrophic membrane failure.")

heading("10. 细胞器接触位点：铁死亡生物学缺失的空间层",1)
para("细胞器接触位点是本文最建议突出为创新中心的章节。传统综述多问“哪个细胞器参与铁死亡”，而更高层次的问题是：细胞器如何在铁死亡过程中交换脂质、Ca²⁺、铁和氧化还原信号？")
bullet("ER–mitochondria contacts：脂质转运、Ca²⁺交换、代谢和ROS整合。")
bullet("ER–peroxisome contacts：ether-lipid和VLCFA相关代谢协同。")
bullet("lysosome–mitochondria contacts：可能参与铁转运、质量控制和ROS放大。")
bullet("lipid-droplet contacts：把脂质储存状态迅速转化为膜脂供给状态。")

heading("11. 铁死亡细胞器失败的时间模型",1)
para("目前没有适用于所有细胞的固定“细胞器损伤顺序”。建议采用三阶段模型：Stage I priming（PUFA-PL重塑、铁积累、抗氧化储备下降）；Stage II amplification（mitochondrial ROS、lysosomal iron mobilization、ER stress和contact-site remodeling）；Stage III execution（广泛膜脂过氧化、质膜损伤、离子失衡和细胞死亡）。")
para("该模型的价值在于强调：不同细胞可以通过不同细胞器入口进入同一个终末铁死亡状态。")

heading("12. 细胞器定义的铁死亡疾病内表型",1)
para("疾病章节不应继续简单按“cancer、AKI、neurodegeneration”逐病种罗列，而应进一步问：该疾病中哪一组细胞器网络主导铁死亡？由此可以建立organelle-defined ferroptosis endotypes。")
bullet("癌症：ER–mitochondria–lipid droplet与抗氧化逃逸。")
bullet("缺血再灌注：mitochondria–membrane ROS burst与lipid oxidation。")
bullet("AKI：mitochondria–lysosome iron mobilization。")
bullet("神经退行性疾病：mitochondria–ER–lysosome redox/quality-control failure。")
bullet("肝病：ER–lipid droplet–peroxisome lipid overload。")
bullet("肠道炎症：plasma membrane–mitochondria–ER与屏障膜脂过氧化。")

add_figure(f5,"Figure 5 | Organelle-defined ferroptosis endotypes across major disease settings",
"Different disease settings engage distinct combinations of organelles. The endotype concept shifts the review from disease cataloguing to mechanistic stratification, helping identify which organelle dependency is most relevant to a given cell type, stage and therapeutic goal.")

heading("13. 治疗：从单分子抑制升级到细胞器网络重编程",1)
para("治疗部分必须区分癌症中的ferroptosis induction与缺血、退行性和炎症性损伤中的ferroptosis inhibition。真正有价值的转化框架是按干预层级组织，而不是简单罗列药物。")
bullet("Iron-handling layer：铁螯合、TFRC、ferritin和NCOA4/ferritinophagy。")
bullet("Antioxidant layer：GPX4–GSH、FSP1–CoQ、DHODH–CoQH₂、GCH1–BH4。")
bullet("Lipid-remodeling layer：ACSL4、LPCAT3、SCD1、MUFA/PUFA balance与ether-lipid metabolism。")
bullet("Organelle-stress layer：ER stress、mitochondrial quality control、lysosome/autophagy和peroxisomal metabolism。")
bullet("Terminal membrane layer：radical-trapping antioxidants与membrane-repair mechanisms。")

add_figure(f6,"Figure 6 | Therapeutic targeting of organelle crosstalk in ferroptosis",
"Therapeutic strategies can be organized by the organelle processes that create or suppress ferroptotic stress. Cancer treatment may require selective induction of ferroptosis, whereas tissue injury and neurodegeneration generally require ferroptosis inhibition. Translation will depend on cell-state biomarkers, organelle-specific dependencies, targeted delivery and disease-specific therapeutic windows.")

heading("14. 临床标志物：从‘铁死亡marker’到多维证据组合",1)
para("目前不存在单一指标能够在患者组织中可靠证明铁死亡。建议采用“lipid peroxidation + iron state + defense system + organelle state + spatial information”的组合证据。")
bullet("Lipid peroxidation：oxidized phospholipids、4-HNE、MDA、F₂-isoprostanes。")
bullet("Iron：labile Fe²⁺、ferritin、TFRC及organelle-resolved iron probes。")
bullet("Defense：GPX4、FSP1、DHODH、GSH和CoQ redox state。")
bullet("Organelle state：mitochondrial morphology、ER stress、NCOA4、lysosomal function、peroxisomes和lipid-droplet dynamics。")

heading("15. Outstanding questions",1)
bullet("Which organelle initiates ferroptosis in each cellular context?")
bullet("How are oxidized lipids transferred between organelles?")
bullet("Are organelle contact sites causal nodes or secondary adaptations?")
bullet("What determines cell-type-specific organelle dependencies?")
bullet("Can ferroptosis be targeted without disrupting essential organelle functions?")
bullet("Can organelle-resolved biomarkers be validated prospectively in human disease?")

# Tables
doc.add_page_break()
heading("16. 可编辑综合表格",1)
add_table("Table 1 | Organelle-specific functions in ferroptosis",
["Organelle","Homeostatic role","Pro-ferroptotic mechanisms","Anti-ferroptotic mechanisms","Representative readouts"],[
["Mitochondria","TCA/ETC, Fe–S clusters, metabolic redox","ROS, iron dysregulation, metabolic amplification","DHODH–CoQH₂, metabolic adaptation","mitoROS, ΔΨm, cristae, DHODH"],
["ER","Lipid synthesis, Ca²⁺ storage, proteostasis","ACSL4/LPCAT3 PUFA-PL remodeling, ER stress","adaptive UPR, lipid remodeling","ACSL4, LPCAT3, PERK/ATF4, ER Ca²⁺"],
["Lysosome/autophagosome","Recycling, ferritin/lipid turnover","NCOA4 ferritinophagy, lipophagy, iron release","controlled flux, iron sequestration","NCOA4, LC3/p62, lysosomal Fe²⁺, LAMP1"],
["Peroxisome","VLCFA oxidation, ether-lipid synthesis","PUFA ether-PL synthesis, oxidative metabolism","context-dependent redox/lipid control","AGPS, FAR1, PEX proteins, ether lipidomics"],
["Lipid droplets","TAG storage, FA buffering","lipolysis/lipophagy supplies PUFAs","sequestration of PUFA away from membranes","PLIN2/3, ATGL, LD number/size"],
["Nucleus","Transcription and stress integration","pro-ferroptotic transcriptional programs","NRF2-dependent adaptation","NRF2, p53, ATF4, target genes"],
["Golgi/plasma membrane","Trafficking and membrane integrity","lipid-peroxide propagation, ion imbalance","FSP1–CoQ, system Xc⁻, membrane repair","FSP1, CoQ, SLC7A11, PI/LDH"],
])

add_table("Table 2 | Key molecular regulators linking organelles to ferroptosis",
["Regulatory axis","Principal compartment","Core function","Effect","Interpretation"],[
["GPX4–GSH","multiple membranes/cytosol","reduces phospholipid hydroperoxides","Suppresses","central but not sole defense"],
["FSP1–CoQ","plasma membrane-associated","regenerates reduced CoQ","Suppresses","GPX4-independent membrane defense"],
["DHODH–CoQH₂","mitochondrial inner membrane","mitochondrial CoQ reduction","Suppresses","organelle-specific antioxidant defense"],
["GCH1–BH4","cytosolic metabolic network","radical trapping/lipid remodeling","Suppresses","parallel defense axis"],
["ACSL4–LPCAT3","ER-associated lipid metabolism","PUFA incorporation into phospholipids","Promotes","controls substrate availability"],
["SLC7A11","plasma membrane","cystine uptake for GSH synthesis","Suppresses","upstream GPX4 support"],
["NCOA4 ferritinophagy","lysosome/autophagosome","mobilizes ferritin iron","Often promotes","highly context-dependent"],
["NRF2","nucleus","antioxidant/iron/lipid transcription","Usually suppresses","can cause tumour resistance"],
["PUFA ether-lipid synthesis","peroxisome–ER","generates oxidizable ether phospholipids","Can promote","lipid species/context matter"],
])

add_table("Table 3 | Disease relevance of organelle crosstalk in ferroptosis",
["Disease context","Dominant organelle axes","Core mechanism","Potential biomarkers","Therapeutic direction"],[
["Cancer","ER–mitochondria–lipid droplet","PUFA remodeling, metabolic stress, antioxidant escape","ACSL4, GPX4, FSP1, lipidomics","Induce ferroptosis"],
["Ischaemia–reperfusion","mitochondria–membrane/peroxisome","ROS burst and lipid oxidation","4-HNE, MDA, mitochondrial injury","Inhibit ferroptosis"],
["Acute kidney injury","mitochondria–lysosome","iron release, mitochondrial ROS","NGAL/KIM-1 + ferroptosis panel","Inhibit ferroptosis"],
["Neurodegeneration","mitochondria–ER–lysosome","ROS, ER stress, iron accumulation","lipid peroxide + iron/imaging markers","Inhibit ferroptosis"],
["Liver disease","ER–lipid droplet–peroxisome","lipid overload and PUFA remodeling","hepatic lipidomics, 4-HNE, iron","Context-dependent inhibition"],
["Intestinal injury/IBD","plasma membrane–mitochondria–ER","epithelial lipid peroxidation and barrier injury","barrier + lipid-peroxidation markers","Protect epithelium/inhibit ferroptosis"],
])

add_table("Table 4 | Outstanding questions and future directions",
["Theme","Current gap","Why it matters","Recommended tools","Priority"],[
["Organelle contact sites","causal roles incompletely defined","identifies spatial signaling nodes","proximity labeling, super-resolution, synthetic tethers","Very high"],
["Temporal sequence","initiating organelle is unclear","defines early intervention window","live-cell multi-organelle imaging","Very high"],
["Compartment lipid peroxidation","lethal lipid species poorly resolved","distinguishes local functions","subcellular lipidomics, MSI, organelle probes","Very high"],
["Cell-type specificity","organelle dependencies differ","enables precision treatment","single-cell/spatial multi-omics","High"],
["Selective autophagy","protective vs pro-ferroptotic switch unclear","prevents overgeneralization","cargo-specific flux/perturbation","High"],
["Disease heterogeneity","endotypes not prospectively defined","supports stratification","clinical cohorts, spatial omics","High"],
["Organelle-targeted therapeutics","delivery remains difficult","may improve efficacy/safety","organelle-targeted nanoparticles/prodrugs","High"],
["Clinical biomarkers","no validated single marker","required for patient trials","multimodal panels","Very high"],
])

heading("17. Perspective：这篇综述真正需要强调的四个创新支点",1)
bullet("Organelle contact sites：从‘细胞器参与’升级到‘细胞器如何交流’。")
bullet("Spatiotemporal sequence：从终点分子变化升级到铁死亡发生的时间顺序。")
bullet("Compartment-specific lipid peroxidation：从全细胞ROS升级到区室特异性氧化脂质。")
bullet("Organelle-defined disease endotypes：从病名分类升级到机制依赖分类。")
para("如果全文始终围绕这四个支点组织，就能明显区别于一般的“线粒体一节、ER一节、溶酶体一节”的综述。")

heading("18. Conclusion",1)
para("铁死亡并非单一氧化反应的终点，而是多个细胞器围绕铁稳态、脂质重塑、氧化还原防御和膜系统完整性进行协同与失衡的结果。线粒体决定关键代谢和氧化环境；ER塑造易过氧化膜脂；过氧化物酶体参与特殊脂质合成；溶酶体与自噬控制铁脂动员；脂滴缓冲或释放脂肪酸；细胞核决定应激适应；Golgi和质膜连接物流与终末膜损伤。")
para("因此，本文建议把铁死亡定义为：a systems-level collapse of organelle-coordinated iron, lipid, redox and membrane homeostasis。未来最重要的方向不是继续扩增调控分子名单，而是解析细胞器接触位点、区室特异性脂质过氧化、时间顺序、细胞类型异质性和临床可检测的细胞器状态，并据此建立精准干预框架。")

heading("19. 关键证据锚点（建议作为正式写作起始文献）",1)
refs=[
"Dixon SJ, Olzmann JA. The cell biology of ferroptosis. Nature Reviews Molecular Cell Biology. 2024;25:424–442.",
"Jiang X, Stockwell BR, Conrad M. Ferroptosis: mechanisms, biology and role in disease. Nature Reviews Molecular Cell Biology. 2021;22:266–282.",
"Zou Y, Henry WS, Ricq EL, et al. Plasticity of ether lipids promotes ferroptosis susceptibility and evasion. Nature. 2020;585:603–608.",
"Bersuker K, Hendricks JM, Li Z, et al. The CoQ oxidoreductase FSP1 acts parallel to GPX4 to inhibit ferroptosis. Nature. 2019;575:688–692.",
"Doll S, Freitas FP, Shah R, et al. FSP1 is a glutathione-independent ferroptosis suppressor. Nature. 2019;575:693–698.",
"Mao C, Liu X, Zhang Y, et al. DHODH-mediated ferroptosis defence is a targetable vulnerability in cancer. Nature. 2021;593:586–590.",
]
for rtxt in refs:
    para(rtxt,size=9.2,indent=False,space=2)

doc.core_properties.title="细胞器与铁死亡串扰：综述框架与图表整合版"
doc.core_properties.subject="Organelle crosstalk in ferroptosis review framework"
doc.core_properties.author="ChatGPT"
doc.save(OUT)

# Structural QA
check=Document(OUT)
assert len(check.inline_shapes) == 6, len(check.inline_shapes)
assert len(check.tables) == 4, len(check.tables)
assert len(check.paragraphs) > 80
assert OUT.stat().st_size > 500000

# DOCX zip integrity
with zipfile.ZipFile(OUT,"r") as z:
    assert z.testzip() is None
    assert "word/document.xml" in z.namelist()

print(str(OUT))
print("paragraphs",len(check.paragraphs),"figures",len(check.inline_shapes),"tables",len(check.tables),"bytes",OUT.stat().st_size)
