
from pathlib import Path
import math
import textwrap
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle, FancyArrowPatch, Rectangle, Ellipse
from matplotlib.lines import Line2D

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTDIR = Path("chatgpt_exports")
OUTDIR.mkdir(exist_ok=True)
FIGDIR = OUTDIR / "pparg_figures"
FIGDIR.mkdir(exist_ok=True)
OUT = OUTDIR / "PPARG_Gut_Related_Diseases_Review_Framework_Figures_Tables_CN.docx"

plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "font.size": 9,
    "figure.dpi": 180,
})

COL = {
    "blue":"#DCEEFF","blue2":"#75AADB","green":"#DDF4E6","green2":"#5FAF79",
    "pink":"#FCE2E7","pink2":"#CF7184","purple":"#E9E2FA","purple2":"#8D79C6",
    "orange":"#FCEBD6","orange2":"#D69A42","teal":"#DDF3F2","teal2":"#4E9D9B",
    "yellow":"#FFF3C7","red":"#F5C4C0","dark":"#17365D","grey":"#F4F6F8"
}

def add_box(ax, xy, w, h, title, lines, face, edge, title_size=10, text_size=8, radius=0.02):
    x,y=xy
    box=FancyBboxPatch((x,y),w,h,boxstyle=f"round,pad=0.012,rounding_size={radius}",
                       linewidth=1.2,edgecolor=edge,facecolor=face)
    ax.add_patch(box)
    ax.text(x+0.02*w, y+h-0.08*h, title, ha="left", va="top",
            fontsize=title_size, fontweight="bold", color=COL["dark"])
    yy=y+h-0.24*h
    for line in lines:
        wrap=textwrap.wrap(line, width=max(18,int(36*w/0.25)))
        for s in wrap:
            ax.text(x+0.03*w, yy, "• "+s, ha="left", va="top", fontsize=text_size, color="#22313F")
            yy-=0.075*h
        yy-=0.012*h
    return box

def arrow(ax, a, b, color="#4A5A6A", lw=1.5, style="-|>", rad=0):
    ax.add_patch(FancyArrowPatch(a,b,arrowstyle=style,mutation_scale=13,
                                 linewidth=lw,color=color,connectionstyle=f"arc3,rad={rad}"))

def save(fig, name):
    path=FIGDIR/name
    fig.savefig(path,bbox_inches="tight",facecolor="white")
    plt.close(fig)
    return path

# ---------------- Figure 1 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 1 | PPARG as a central integrator of intestinal homeostasis and gut-related disease",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")

add_box(ax,(0.02,0.55),0.29,0.34,"1. Epithelial barrier homeostasis",[
    "Goblet-cell mucus production and mucin barrier",
    "Tight-junction integrity: claudins, occludin and ZO-1",
    "Epithelial renewal, wound restitution and Paneth-cell defense",
    "Colonocyte fatty-acid oxidation and physiologic hypoxia"
],COL["green"],COL["green2"],11,8.2)

add_box(ax,(0.69,0.55),0.29,0.34,"2. Microbiota–metabolite regulation",[
    "SCFAs, tryptophan metabolites, secondary bile acids and microbial lipids",
    "Reciprocal feedback between PPARG activity and the luminal ecosystem",
    "Dysbiosis alters ligand availability, oxygenation and inflammatory tone"
],COL["purple"],COL["purple2"],11,8.2)

add_box(ax,(0.02,0.12),0.29,0.34,"3. Immune and stromal control",[
    "Macrophage resolution programs and reduced NF-kB signaling",
    "Tolerogenic dendritic-cell states and Treg/Th17 balance",
    "Restriction of neutrophil-driven tissue injury",
    "Fibroblast quiescence, repair and fibrosis control"
],COL["blue"],COL["blue2"],11,8.2)

add_box(ax,(0.69,0.12),0.29,0.34,"4. Gut-related disease outcomes",[
    "Inflammatory bowel disease and relapsing colitis",
    "Colorectal cancer: context- and stage-dependent effects",
    "Intestinal fibrosis and tissue remodeling",
    "Functional bowel inflammation and gut–liver metabolic disease"
],COL["orange"],COL["orange2"],11,8.2)

# central nucleus
ax.add_patch(Ellipse((0.5,0.50),0.24,0.29,facecolor="#EEF1FF",edgecolor="#7B82B8",lw=1.5))
ax.add_patch(Circle((0.455,0.55),0.045,facecolor=COL["pink"],edgecolor=COL["pink2"],lw=1.2))
ax.add_patch(Circle((0.545,0.55),0.045,facecolor=COL["blue"],edgecolor=COL["blue2"],lw=1.2))
ax.text(0.455,0.55,"PPARγ",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.545,0.55,"RXR",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.5,0.475,"PPREs",ha="center",va="center",fontsize=11,fontweight="bold",color=COL["dark"])
ax.text(0.5,0.43,"Target-gene transcription",ha="center",va="center",fontsize=10)

# ligands
ax.text(0.5,0.865,"Ligands / regulatory inputs",ha="center",fontsize=11,fontweight="bold",color=COL["dark"])
for x,label,c in [(0.39,"fatty acids\n& oxylipins",COL["pink2"]),(0.50,"microbial\nmetabolites",COL["green2"]),(0.61,"drugs /\nselective modulators",COL["purple2"])]:
    ax.add_patch(Circle((x,0.80),0.032,facecolor="white",edgecolor=c,lw=2))
    ax.text(x,0.755,label,ha="center",va="top",fontsize=8)
    arrow(ax,(x,0.765),(0.5,0.64),color=c,lw=1.4)

for a,b,c in [((0.31,0.72),(0.39,0.60),COL["green2"]),((0.69,0.72),(0.61,0.60),COL["purple2"]),
              ((0.31,0.30),(0.39,0.43),COL["blue2"]),((0.69,0.30),(0.61,0.43),COL["orange2"])]:
    arrow(ax,a,b,color=c,lw=2)

ax.add_patch(FancyBboxPatch((0.34,0.02),0.32,0.075,boxstyle="round,pad=0.01,rounding_size=0.02",
                            facecolor="#EAF7EC",edgecolor=COL["green2"],lw=1.3))
ax.text(0.5,0.057,"Balanced PPARG signaling supports barrier,\nimmune, microbial and metabolic homeostasis",
        ha="center",va="center",fontsize=9.5,fontweight="bold",color="#245B3A")
f1=save(fig,"Figure1_PPARG_central_integrator.png")

# ---------------- Figure 2 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 2 | Molecular architecture and regulatory network of PPARγ signaling in the gut",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")
add_box(ax,(0.02,0.55),0.30,0.34,"1. Ligands and activation inputs",[
    "Endogenous fatty acids, eicosanoids and endocannabinoid-related lipids",
    "Microbial SCFAs, indoles, secondary bile acids and microbial lipids",
    "Synthetic agonists, partial agonists and selective PPARG modulators"
],COL["teal"],COL["teal2"],11,8)
add_box(ax,(0.68,0.55),0.30,0.34,"2. Co-regulators and chromatin control",[
    "Co-activators: PGC-1α, SRC-1, CBP/p300 and chromatin-opening complexes",
    "Co-repressors: NCoR, SMRT, HDAC3 and related transcriptional brakes",
    "Cell state and chromatin accessibility determine ligand response"
],COL["blue"],COL["blue2"],11,8)
add_box(ax,(0.02,0.12),0.30,0.34,"3. Post-translational regulation",[
    "Phosphorylation can enhance or restrain receptor activity",
    "Acetylation/deacetylation alters DNA binding and co-regulator recruitment",
    "SUMOylation and ubiquitination shape repression, stability and turnover"
],COL["orange"],COL["orange2"],11,8)
add_box(ax,(0.68,0.12),0.30,0.34,"4. Signaling crosstalk",[
    "Transrepression of NF-κB and AP-1 inflammatory programs",
    "Crosstalk with STAT3, AMPK/SIRT1, HIF and Wnt/β-catenin",
    "The output depends on cell type, metabolic state and inflammatory context"
],COL["purple"],COL["purple2"],11,8)

ax.add_patch(Ellipse((0.5,0.52),0.25,0.32,facecolor="#F1F3FF",edgecolor="#8188BC",lw=1.5))
ax.add_patch(Circle((0.455,0.57),0.047,facecolor=COL["pink"],edgecolor=COL["pink2"],lw=1.2))
ax.add_patch(Circle((0.55,0.57),0.047,facecolor=COL["blue"],edgecolor=COL["blue2"],lw=1.2))
ax.text(0.455,0.57,"PPARγ",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.55,0.57,"RXR",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.50,0.485,"PPREs",ha="center",fontsize=11,fontweight="bold")
ax.text(0.50,0.445,"Context-dependent transcription",ha="center",fontsize=9)
for a,b,c in [((0.32,0.72),(0.40,0.61),COL["teal2"]),((0.68,0.72),(0.60,0.61),COL["blue2"]),
              ((0.32,0.30),(0.40,0.42),COL["orange2"]),((0.68,0.30),(0.60,0.42),COL["purple2"])]:
    arrow(ax,a,b,color=c,lw=2)

# outputs ribbon
outs=["fatty-acid oxidation","mitochondrial fitness","barrier genes","mucin","anti-inflammatory mediators","epithelial repair"]
xs=[0.12,0.27,0.42,0.57,0.72,0.87]
for x,label in zip(xs,outs):
    ax.add_patch(FancyBboxPatch((x-0.065,0.015),0.13,0.07,boxstyle="round,pad=0.008,rounding_size=0.012",
                                facecolor=COL["pink"],edgecolor=COL["pink2"],lw=0.8))
    ax.text(x,0.05,label,ha="center",va="center",fontsize=7.6,fontweight="bold")
arrow(ax,(0.50,0.36),(0.50,0.10),color=COL["pink2"],lw=2)
f2=save(fig,"Figure2_PPARG_regulatory_network.png")

# ---------------- Figure 3 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 3 | PPARγ preserves the intestinal epithelial barrier and metabolic fitness",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")
titles=["1. Mucus barrier","2. Tight junctions","3. Epithelial renewal","4. Metabolic fitness","5. Antimicrobial defense"]
details=[
["Goblet-cell differentiation","MUC2 production","Mucus-layer thickness"],
["Claudins, occludin and ZO-1","Reduced paracellular permeability","Barrier to luminal antigens"],
["Stem/progenitor-cell support","Balanced differentiation","Wound restitution"],
["Fatty-acid β-oxidation","Physiologic epithelial hypoxia","Lower ROS and HIF balance"],
["Paneth-cell function","Defensins and lysozyme","Restriction of bacterial encroachment"]
]
faces=[COL["green"],COL["blue"],COL["purple"],COL["orange"],COL["pink"]]
edges=[COL["green2"],COL["blue2"],COL["purple2"],COL["orange2"],COL["pink2"]]
for i in range(5):
    x=0.015+i*0.196
    add_box(ax,(x,0.59),0.185,0.29,titles[i],details[i],faces[i],edges[i],10.5,7.8)

# central epithelial strip
ax.add_patch(Rectangle((0.06,0.38),0.88,0.12,facecolor="#F9D9C8",edgecolor="#C5856A",lw=1.0))
for i in range(14):
    x=0.075+i*0.063
    ax.add_patch(FancyBboxPatch((x,0.39),0.048,0.10,boxstyle="round,pad=0.004,rounding_size=0.01",
                                facecolor="#FFEBDD",edgecolor="#D79E83",lw=0.7))
ax.text(0.5,0.445,"Intestinal epithelium: PPARG links barrier structure to oxidative metabolism",
        ha="center",va="center",fontsize=10,fontweight="bold",color="#6A3C2D")
# bottom failures
fails=[
("Mucus thinning","↓ mucin and goblet-cell function"),
("Leaky barrier","↓ tight-junction integrity"),
("Impaired repair","↓ renewal and restitution"),
("Metabolic stress","↓ β-oxidation; ↑ ROS/O₂"),
("Microbial encroachment","↓ antimicrobial defense")
]
for i,(t,d) in enumerate(fails):
    x=0.015+i*0.196
    ax.add_patch(FancyBboxPatch((x,0.10),0.185,0.18,boxstyle="round,pad=0.01,rounding_size=0.012",
                                facecolor="#FDEBEC",edgecolor="#D47A7A",lw=0.9))
    ax.text(x+0.01,0.245,t,ha="left",va="top",fontsize=9.2,fontweight="bold",color="#9E2F2F")
    ax.text(x+0.01,0.205,d,ha="left",va="top",fontsize=7.6,color="#4B3030")
    arrow(ax,(x+0.092,0.37),(x+0.092,0.29),color="#C64A4A",lw=1.2)
ax.add_patch(FancyBboxPatch((0.14,0.02),0.72,0.055,boxstyle="round,pad=0.01,rounding_size=0.015",
                            facecolor="#E9F8EC",edgecolor=COL["green2"],lw=1.2))
ax.text(0.5,0.047,"Beneficial PPARγ signaling preserves the barrier, supports tissue recovery and constrains dysbiosis",
        ha="center",va="center",fontsize=10,fontweight="bold",color="#245B3A")
f3=save(fig,"Figure3_PPARG_barrier_metabolic_fitness.png")

# ---------------- Figure 4 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 4 | The microbiota–metabolite–PPARγ axis in intestinal homeostasis and disease",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")
add_box(ax,(0.02,0.48),0.29,0.39,"Eubiosis / homeostasis",[
    "Diverse commensals generate SCFAs, indoles, secondary bile acids and microbial lipids",
    "Robust mucus and tight-junction barrier",
    "Low luminal oxygen and strong colonization resistance"
],COL["green"],COL["green2"],11,8.2)
add_box(ax,(0.69,0.48),0.29,0.39,"Dysbiosis / disease",[
    "Reduced beneficial metabolites and increased inflammatory microbial products",
    "Barrier leak and higher epithelial oxygenation",
    "Pathobiont expansion and chronic inflammatory signaling"
],COL["pink"],COL["pink2"],11,8.2)

# metabolite row
metabs=[("SCFAs","#64B66E"),("Indoles","#8E79C8"),("Secondary\nbile acids","#E0A542"),("Microbial\nlipids","#559BC9")]
for i,(lab,c) in enumerate(metabs):
    x=0.37+i*0.087
    ax.add_patch(Circle((x,0.78),0.027,facecolor=c,edgecolor="white",lw=1))
    ax.text(x,0.735,lab,ha="center",va="top",fontsize=7.8)
    arrow(ax,(x,0.75),(0.5,0.64),color=c,lw=1.5)

ax.add_patch(Ellipse((0.5,0.52),0.24,0.25,facecolor="#EEF2FF",edgecolor="#7B82B8",lw=1.3))
ax.add_patch(Circle((0.47,0.55),0.05,facecolor=COL["pink"],edgecolor=COL["pink2"]))
ax.text(0.47,0.55,"PPARγ",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.56,0.55,"RXR",ha="center",va="center",fontsize=10,fontweight="bold",color=COL["dark"])
ax.text(0.50,0.47,"Host transcriptional\nresponse",ha="center",va="center",fontsize=9)

outlabs=["β-oxidation","barrier support","epithelial hypoxia","colonization resistance","anti-inflammatory tone"]
for i,lab in enumerate(outlabs):
    x=0.25+i*0.125
    ax.add_patch(FancyBboxPatch((x-0.052,0.27),0.104,0.075,boxstyle="round,pad=0.008,rounding_size=0.012",
                                facecolor=COL["blue"],edgecolor=COL["blue2"],lw=0.9))
    ax.text(x,0.307,lab,ha="center",va="center",fontsize=7.3,fontweight="bold")
    arrow(ax,(0.5,0.40),(x,0.35),color="#4F82B7",lw=1.0)
arrow(ax,(0.31,0.62),(0.39,0.56),color=COL["green2"],lw=2)
arrow(ax,(0.69,0.62),(0.61,0.56),color=COL["pink2"],lw=2)
ax.text(0.18,0.16,"Reciprocal feedback:\nPPARG shapes the mucosal environment\nthat selects for commensal functions",
        ha="center",va="center",fontsize=9,color="#2E6650")
ax.text(0.82,0.16,"Failure of the axis:\nreduced metabolites + barrier leak\n→ inflammatory amplification",
        ha="center",va="center",fontsize=9,color="#963B47")
arrow(ax,(0.34,0.20),(0.42,0.28),color=COL["green2"],lw=1.5,rad=-0.1)
arrow(ax,(0.66,0.20),(0.58,0.28),color=COL["pink2"],lw=1.5,rad=0.1)
f4=save(fig,"Figure4_Microbiota_metabolite_PPARG_axis.png")

# ---------------- Figure 5 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 5 | Immune and stromal actions of PPARγ in gut inflammation",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")
modules=[
("Macrophages",["M1→pro-resolving shift","↑ IL-10 / efferocytosis","↓ TNF, IL-1β and ROS"],COL["blue"],COL["blue2"]),
("Dendritic cells",["Tolerogenic phenotype","↓ antigen presentation","Support Treg induction"],COL["purple"],COL["purple2"]),
("Treg / Th17",["↑ Treg / FOXP3 / IL-10","↓ Th17 / IL-17","Immune-tolerance bias"],COL["orange"],COL["orange2"]),
("Neutrophils",["↓ recruitment and degranulation","↓ NET-related tissue injury","Faster resolution"],COL["pink"],COL["pink2"]),
("Fibroblasts",["↓ myofibroblast activation","↓ ECM deposition","Support repair without fibrosis"],COL["teal"],COL["teal2"]),
("Epithelial–immune crosstalk",["↑ barrier and mucus","↓ epithelial cytokines","Reduced immune activation"],COL["green"],COL["green2"])
]
positions=[(0.02,0.58),(0.68,0.58),(0.02,0.30),(0.68,0.30),(0.02,0.02),(0.68,0.02)]
for (title,lines,face,edge),(x,y) in zip(modules,positions):
    add_box(ax,(x,y),0.30,0.22,title,lines,face,edge,10.5,7.7)
# central
ax.add_patch(Ellipse((0.5,0.48),0.27,0.32,facecolor="#F1F3FF",edgecolor="#8188BC",lw=1.5))
ax.add_patch(Circle((0.46,0.54),0.05,facecolor=COL["pink"],edgecolor=COL["pink2"],lw=1.2))
ax.add_patch(Circle((0.56,0.54),0.05,facecolor=COL["blue"],edgecolor=COL["blue2"],lw=1.2))
ax.text(0.46,0.54,"PPARγ",ha="center",va="center",fontsize=12,fontweight="bold")
ax.text(0.56,0.54,"RXR",ha="center",va="center",fontsize=11,fontweight="bold")
ax.text(0.51,0.445,"Immune-resolution\ntranscriptional program",ha="center",va="center",fontsize=9)
for x,y in [(0.32,0.69),(0.68,0.69),(0.32,0.41),(0.68,0.41),(0.32,0.13),(0.68,0.13)]:
    arrow(ax,(0.50,0.48),(x,y),color="#73849B",lw=1.5)
ax.add_patch(FancyBboxPatch((0.35,0.12),0.30,0.16,boxstyle="round,pad=0.012,rounding_size=0.02",
                            facecolor="#EAF7EC",edgecolor=COL["green2"],lw=1.2))
ax.text(0.50,0.23,"Net mediator shift",ha="center",fontsize=10,fontweight="bold",color=COL["dark"])
ax.text(0.50,0.19,"↓ TNF / IL-1β / IL-6 / CXCL8 / ROS",ha="center",fontsize=8.5,color="#8A3232")
ax.text(0.50,0.155,"↑ IL-10 / TGF-β / pro-resolving programs",ha="center",fontsize=8.5,color="#256443")
f5=save(fig,"Figure5_Immune_stromal_PPARG.png")

# ---------------- Figure 6 ----------------
fig,ax=plt.subplots(figsize=(12.5,8.2)); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis("off")
ax.text(0.02,0.97,"Figure 6 | Therapeutic targeting of the PPARγ axis in gut-related disease",
        fontsize=16,fontweight="bold",color=COL["dark"],va="top")
layers=[
("1. Dietary modulation",["Mediterranean-style pattern","fiber / prebiotics","n-3 PUFAs / polyphenols"],COL["green"],COL["green2"]),
("2. Microbiota-directed",["prebiotics / probiotics","postbiotic metabolites","FMT: context-specific"],COL["blue"],COL["blue2"]),
("3. Endogenous metabolites",["SCFAs / indoles","secondary bile acids","pro-resolving lipids"],COL["orange"],COL["orange2"]),
("4. Direct agonists",["TZDs: rosiglitazone / pioglitazone","repurposing opportunities","systemic adverse effects"],COL["purple"],COL["purple2"]),
("5. Selective modulators",["partial agonists / SPPARMs","tissue-selective activity","improved safety goal"],COL["teal"],COL["teal2"]),
("6. Combination therapy",["with anti-inflammatory agents","with anti-fibrotics","metabolic co-targeting"],COL["pink"],COL["pink2"]),
("7. Intestine-targeted delivery",["colon-targeted formulations","nanoparticles / hydrogels","microbiota-responsive release"],COL["yellow"],COL["orange2"])
]
for i,(title,lines,face,edge) in enumerate(layers):
    x=0.012+i*0.141
    add_box(ax,(x,0.56),0.132,0.31,title,lines,face,edge,9.2,7.0)

# disease mapping
diseases=[
("IBD / colitis","barrier + inflammation"),
("CRC","context-dependent tumor biology"),
("intestinal fibrosis","myofibroblast / ECM"),
("IBS","low-grade inflammation"),
("gut–liver disease","metabolism + endotoxemia")
]
for i,(d,desc) in enumerate(diseases):
    x=0.03+i*0.19
    ax.add_patch(FancyBboxPatch((x,0.31),0.17,0.14,boxstyle="round,pad=0.01,rounding_size=0.015",
                                facecolor="#FFF5F1",edgecolor="#D99682",lw=1))
    ax.text(x+0.085,0.405,d,ha="center",va="center",fontsize=9,fontweight="bold",color=COL["dark"])
    ax.text(x+0.085,0.35,desc,ha="center",va="center",fontsize=7.5,color="#5A4C49")
    arrow(ax,(x+0.085,0.55),(x+0.085,0.46),color="#9A8B74",lw=1)
# translational considerations
trans=[
("Biomarkers","PPARG activity, barrier markers,\nmetabolites, inflammatory signatures"),
("Patient stratification","disease subtype, microbiome,\nmetabolic phenotype"),
("Targeted delivery","colon-specific exposure,\nminimal systemic exposure"),
("Efficacy endpoints","mucosal healing, barrier,\ninflammation, fibrosis"),
("Safety","weight gain/edema,\ncardiometabolic and tumor context")
]
for i,(t,d) in enumerate(trans):
    x=0.03+i*0.19
    ax.add_patch(FancyBboxPatch((x,0.08),0.17,0.15,boxstyle="round,pad=0.01,rounding_size=0.015",
                                facecolor=COL["grey"],edgecolor="#AAB3BD",lw=0.9))
    ax.text(x+0.085,0.195,t,ha="center",fontsize=8.8,fontweight="bold",color=COL["dark"])
    ax.text(x+0.085,0.145,d,ha="center",va="center",fontsize=7.0)
ax.add_patch(FancyBboxPatch((0.12,0.015),0.76,0.045,boxstyle="round,pad=0.008,rounding_size=0.012",
                            facecolor="#EAF7EC",edgecolor=COL["green2"],lw=1.2))
ax.text(0.5,0.037,"Precision modulation—not indiscriminate activation—should restore local barrier, immune and metabolic homeostasis",
        ha="center",va="center",fontsize=9.3,fontweight="bold",color="#245B3A")
f6=save(fig,"Figure6_Therapeutic_targeting_PPARG.png")

# ---------------- DOCX helpers ----------------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_font(run, size=10.5, bold=False, color=None, east="宋体", latin="Times New Roman"):
    run.font.name=latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size=Pt(size)
    run.bold=bold
    if color:
        run.font.color.rgb=RGBColor.from_string(color)

def p(doc,text="",size=10.5,bold=False,align=None,indent=True,after=5):
    para=doc.add_paragraph()
    if align is not None: para.alignment=align
    para.paragraph_format.space_after=Pt(after)
    if indent: para.paragraph_format.first_line_indent=Cm(0.74)
    r=para.add_run(text); set_font(r,size,bold)
    return para

def h(doc,text,level=1):
    para=doc.add_paragraph()
    para.paragraph_format.space_before=Pt(10 if level==1 else 6)
    para.paragraph_format.space_after=Pt(5)
    r=para.add_run(text)
    if level==1: set_font(r,15,True,"17365D","黑体")
    elif level==2: set_font(r,12.5,True,"245B78","黑体")
    else: set_font(r,11,True,"3B6A57","黑体")
    return para

def bullets(doc,items):
    for item in items:
        para=doc.add_paragraph()
        para.paragraph_format.left_indent=Cm(0.55)
        para.paragraph_format.first_line_indent=Cm(-0.3)
        para.paragraph_format.space_after=Pt(2)
        r=para.add_run("• "+item); set_font(r,10.2)

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers))
    t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    set_repeat_table_header(t.rows[0])
    for j,hh in enumerate(headers):
        c=t.rows[0].cells[j]; set_cell_shading(c,"DCEAF5")
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        rp=c.paragraphs[0]; rp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        rr=rp.add_run(hh); set_font(rr,8.7,True,"17365D","黑体")
    for row in rows:
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pp=cells[j].paragraphs[0]; pp.paragraph_format.space_after=Pt(0)
            rr=pp.add_run(str(val)); set_font(rr,8.1)
    doc.add_paragraph()
    return t

def add_fig(doc,path,title,legend):
    doc.add_picture(str(path),width=Inches(6.55))
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    pp=doc.add_paragraph(); pp.paragraph_format.space_before=Pt(3); pp.paragraph_format.space_after=Pt(2)
    rr=pp.add_run(title); set_font(rr,9.8,True,"17365D")
    pp2=doc.add_paragraph(); pp2.paragraph_format.space_after=Pt(8)
    rr2=pp2.add_run(legend); set_font(rr2,9.0)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
style=doc.styles["Normal"]; style.font.name="Times New Roman"; style._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体"); style.font.size=Pt(10.5)

# cover
pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=pp.add_run("PPARG通路与肠道相关疾病"); set_font(rr,20,True,"17365D","黑体")
pp=doc.add_paragraph(); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=pp.add_run("综述框架、机制图与表格整合版"); set_font(rr,14,True,"245B78","黑体")
p(doc,"定位：影响因子 >10；面向 Nature Reviews / Cell / CNS 子刊风格综述。",10.5,False,WD_ALIGN_PARAGRAPH.CENTER,False,5)
p(doc,"推荐英文题目：PPARγ in the intestinal ecosystem: linking epithelial metabolism, host–microbiota crosstalk and gut-related disease",10.3,False,WD_ALIGN_PARAGRAPH.CENTER,False,4)
doc.add_page_break()

h(doc,"1. 综述定位与中心科学问题",1)
p(doc,"建议避免把本文写成“PPARG的基础知识 + IBD/CRC/NAFLD逐病罗列”。更有竞争力的写法，是把PPARγ定义为连接肠上皮氧化代谢、屏障完整性、微生物生态、免疫耐受与组织修复的“intestinal ecosystem integrator”。这样可以形成一条具有解释力的主线，而不是文献堆砌。")
bullets(doc,[
    "核心主线：ligand/metabolite inputs → PPARγ–RXR transcriptional control → epithelial metabolic fitness → barrier and oxygenation → microbiota ecology → immune/stromal resolution → disease-specific outcomes。",
    "核心悖论：PPARγ通常被视为抗炎/代谢保护因子，但在不同细胞、疾病阶段和肿瘤背景下，其作用具有明显context dependence。",
    "核心转化问题：临床真正需要的可能不是“全身强激动PPARγ”，而是cell-selective、partial或intestine-targeted modulation。",
    "核心证据原则：必须把直接PPARγ因果证据与泛化的“抗炎、代谢改善或微生物变化”分开，避免把所有效应都归因于PPARG。"
])

h(doc,"2. 建议题目",1)
bullets(doc,[
    "首选：PPARγ in the intestinal ecosystem: linking epithelial metabolism, host–microbiota crosstalk and gut-related disease",
    "备选：PPARγ at the gut interface: epithelial metabolism, immune tolerance and microbiota-dependent disease",
    "备选：Context-dependent PPARγ signaling in gut health and disease: from barrier metabolism to precision therapeutics",
    "中文：PPARγ通路与肠道相关疾病：从上皮代谢稳态、菌群互作到精准治疗"
])

h(doc,"3. 推荐正文框架",1)
outline=[
("3.1 Introduction: why PPARγ is more than a metabolic nuclear receptor",[
"从传统脂质/糖代谢受体切入，迅速转向肠道中的屏障、氧化代谢、免疫和微生物生态。",
"指出PPARγ在上皮、髓系细胞、淋巴细胞和基质细胞中的作用不完全相同。",
"提出全文中心模型：PPARγ决定肠道是否维持低氧、低炎症、共生状态。"]),
("3.2 Molecular architecture and regulation of PPARγ signaling",[
"PPARγ–RXR异二聚体、PPRE结合及转录激活/转录抑制。",
"内源性脂质、eicosanoids、endocannabinoids、微生物代谢物和药理学配体。",
"PGC-1α/SRC-1/CBP-p300与NCoR/SMRT/HDAC3等共调节因子。",
"phosphorylation、acetylation、SUMOylation、ubiquitination及染色质情境。"]),
("3.3 PPARγ as a metabolic gatekeeper of the intestinal epithelium",[
"colonocyte fatty-acid β-oxidation、线粒体适能与上皮耗氧。",
"physiologic hypoxia如何限制腔内氧扩散并支持专性厌氧共生菌。",
"tight junctions、mucus、epithelial renewal与Paneth-cell antimicrobial defense。",
"把‘metabolism → oxygenation → ecology’写成文章的第一条机制链。"]),
("3.4 Microbiota–metabolite–PPARγ reciprocal loop",[
"SCFAs、tryptophan-derived indoles、secondary bile acids和microbial lipids作为候选调节因子。",
"PPARγ反过来通过营养/氧环境、黏液、抗菌肽和免疫基线塑造微生物群落。",
"建立从association到causality的证据阶梯：菌群变化 → 代谢物 → receptor engagement → genetic/pharmacologic perturbation → rescue。"]),
("3.5 Immune and stromal PPARγ",[
"macrophage inflammatory-to-resolving reprogramming与efferocytosis。",
"dendritic-cell tolerance与Treg induction；Treg/Th17平衡。",
"neutrophil recruitment/NET-related injury；fibroblast/myofibroblast activation与ECM deposition。",
"强调epithelial PPARγ与immune PPARγ的相互作用，而不是把免疫效应孤立讨论。"]),
("3.6 Disease modules",[
"IBD：UC与CD分别讨论，强调上皮PPARγ、药理学激动和5-ASA相关机制。",
"CRC：把PPARγ的抗肿瘤/促肿瘤争议作为context-dependence的重点章节。",
"intestinal fibrosis：myofibroblast、TGF-β/SMAD与代谢状态。",
"IBS/functional gut inflammation：证据相对新兴，谨慎讨论低度炎症、屏障和gut–brain axis。",
"infectious colitis：讨论过度炎症控制与宿主防御之间的平衡。",
"gut–liver metabolic disease：肠屏障、LPS/endotoxemia、胆汁酸和代谢互作。"]),
("3.7 Therapeutic translation",[
"饮食与microbiota-directed approaches、endogenous metabolite support。",
"TZDs的再利用价值与系统性副作用。",
"partial agonists/SPPARMs、dual/pan-PPAR ligands与biased modulation。",
"intestine-targeted delivery、microbiota-responsive formulations和combination therapy。",
"明确疗效终点：mucosal healing、barrier biomarkers、metabolite signatures、fibrosis endpoints和patient-reported outcomes。"]),
("3.8 Outstanding questions",[
"cell-type specificity、microbiota–metabolite causality、epithelial versus immune PPARγ。",
"CRC中的保护性/促肿瘤作用何时切换。",
"缺乏可用于患者分层的PPARγ activation biomarkers。",
"如何实现肠道局部靶向并降低全身代谢、骨、心血管和潜在肿瘤风险。"])
]
for title,items in outline:
    h(doc,title,2); bullets(doc,items)

h(doc,"4. 适合CNS子刊的差异化写作策略",1)
bullets(doc,[
    "把“PPARγ–microbiota”从菌属变化提升到功能层：substrate → metabolite → receptor → cell state → tissue phenotype。",
    "把“PPARγ抗炎”拆成至少四种机制：transrepression、metabolic rewiring、barrier restoration、resolution/efferocytosis。",
    "IBD、CRC、fibrosis和gut–liver disease不要使用同一个治疗叙事；PPARγ在不同疾病中的风险–收益窗口不同。",
    "正文持续区分parent receptor biology、drug pharmacology和microbiome-mediated indirect effects。",
    "建议在全文末尾设置Perspective/Outstanding Questions，以cell specificity、causality、biomarkers和targeted delivery收束。"
])

doc.add_page_break()
h(doc,"5. 机制图体系",1)
figs=[
(f1,"Figure 1 | PPARγ as a central integrator of intestinal homeostasis and gut-related disease",
"总览图。将PPARγ置于上皮屏障、微生物–代谢物、免疫/基质控制和疾病结局四个模块中央，突出“intestinal ecosystem integrator”概念。"),
(f2,"Figure 2 | Molecular architecture and regulatory network of PPARγ signaling in the gut",
"展示内源性、微生物及药理学配体，共激活/共抑制因子，翻译后修饰及NF-κB、AP-1、STAT3、AMPK/SIRT1、HIF与Wnt/β-catenin等交叉信号。"),
(f3,"Figure 3 | PPARγ preserves the intestinal epithelial barrier and metabolic fitness",
"围绕mucus、tight junction、epithelial renewal、β-oxidation/physiologic hypoxia和Paneth-cell antimicrobial defense五个功能模块，对比PPARγ受损后的屏障失败。"),
(f4,"Figure 4 | The microbiota–metabolite–PPARγ axis in intestinal homeostasis and disease",
"把eubiosis与dysbiosis放在同一张机制图中，通过SCFAs、indoles、secondary bile acids和microbial lipids连接微生物功能与宿主PPARγ，并强调双向反馈和因果链。"),
(f5,"Figure 5 | Immune and stromal actions of PPARγ in gut inflammation",
"整合macrophages、dendritic cells、Treg/Th17、neutrophils、fibroblasts和epithelial–immune crosstalk，突出PPARγ对炎症消退、免疫耐受和纤维化的细胞特异性控制。"),
(f6,"Figure 6 | Therapeutic targeting of the PPARγ axis in gut-related disease",
"从diet、microbiota、endogenous metabolites、direct agonists、selective modulators、combination therapy到intestine-targeted delivery分层展示治疗路线，并加入患者分层、疗效终点与安全性。")
]
for path,title,legend in figs:
    add_fig(doc,path,title,legend)

doc.add_page_break()
h(doc,"6. 原生可编辑Word表格",1)

h(doc,"Table 1 | PPARγ in major gut-related diseases",2)
table(doc,["疾病","主要细胞情境","PPARγ状态/模式","主要机制联系","转化意义","证据成熟度"],[
["Ulcerative colitis","colonocytes、lamina propria macrophages、DCs","炎症黏膜中常见降低/功能不足","屏障修复不足；促炎因子增加；microbiota–host crosstalk异常","局部PPARγ激活、5-ASA相关机制、促进mucosal healing","中高"],
["Crohn’s disease","epithelial cells、monocytes/macrophages、Th1/Th17","病灶区域常呈降低或失衡","屏障破坏、微生物转位、Th1/Th17炎症、消退不足","PPARγ激动/调节可能作为联合策略，需按亚型评估","中等"],
["Colorectal cancer","tumor epithelium、TAMs、CAFs","高度context-和stage-dependent","分化、增殖、凋亡、Wnt/β-catenin、代谢和免疫微环境","必须按肿瘤PPARγ状态和分期分层，避免简单定义为抗癌靶点","中等"],
["Intestinal fibrosis","myofibroblasts、epithelium、macrophages","纤维化组织中可能下降/失衡","TGF-β/SMAD、ECM沉积、代谢和炎症互作","选择性激动可能限制myofibroblast activation和狭窄进展","中等"],
["IBS / functional bowel inflammation","enteric neurons、epithelium、immune/stromal cells","altered / context-dependent","低度炎症、visceral hypersensitivity、屏障和microbiota异常","更适合定义为新兴机制而非成熟治疗靶点","新兴"],
["Infectious colitis","epithelium、macrophages、DCs、neutrophils","急性感染期可出现动态下降/重编程","宿主防御、屏障保护与组织损伤之间的平衡","需要避免因过度激活PPARγ而削弱病原清除","新兴/多为前临床"],
["Gut–liver metabolic disease","IEC、hepatocytes、Kupffer cells、microbiota","肠肝组织呈组织特异性失衡","lipid metabolism、insulin signaling、LPS/endotoxemia、bile acids","有机会通过gut–liver axis和局部/系统联合干预","中等"]
])

h(doc,"Table 2 | Endogenous, microbial and pharmacological modulators of PPARγ relevant to the gut",2)
table(doc,["调节因子类别","代表例子","来源","主要靶细胞","肠道主要效应","转化备注"],[
["Endogenous fatty acids","oleic/linoleic/arachidonic acids、n-3 PUFAs","diet / host lipid metabolism","IEC、macrophage、DC、stroma","PPARγ activation、barrier support、metabolic homeostasis","适合作为营养–受体连接层，但体内受多条脂质通路共同影响"],
["Eicosanoids","15d-PGJ2等","host prostaglandin metabolism","IEC、macrophage、DC、T cells","anti-inflammatory / pro-resolving signals","浓度、时相和局部生成量决定效应"],
["Endocannabinoid-related lipids","2-AG、AEA","host lipid signaling","IEC、immune cells","PPARγ与cannabinoid receptor交叉；barrier/immune effects","存在明显多靶点性"],
["SCFAs","butyrate、propionate、acetate","microbial fiber fermentation","IEC、macrophage、DC、T cells","barrier、regulatory immunity、metabolism","PPARγ只是其多个宿主靶点之一"],
["Tryptophan-derived indoles","IPA、IAA、IAld等","microbial tryptophan metabolism","IEC、immune cells","barrier、antioxidant、immune tolerance","需与AhR/PXR等受体效应区分"],
["Secondary bile acids","DCA、LCA及衍生物","microbial bile-acid transformation","IEC、immune/stromal cells","metabolism、barrier、immune modulation","常与FXR/TGR5通路并行或交叉"],
["TZDs","rosiglitazone、pioglitazone","pharmacological","IEC、immune/stromal cells","potent receptor activation；anti-inflammatory / metabolic effects","全身副作用限制长期肠病应用"],
["Selective modulators","SPPARMs、partial/dual agonists","next-generation pharmacology","cell/tissue dependent","希望保留barrier/anti-inflammatory效应并减少副作用","关键方向：partial、biased、local delivery"]
])

h(doc,"Table 3 | Cell-type-specific functions of PPARγ in the intestinal ecosystem",2)
table(doc,["细胞类型","生理作用","疾病相关失调","关键下游程序","建议读出","治疗意义"],[
["Absorptive epithelial cells / colonocytes","barrier、FAO、renewal、anti-inflammatory tone","IBD中表达/活性降低；氧化代谢下降","tight junction、FAO、NF-κB repression、restitution","Claudin/Occludin/ZO-1、CPT1A/ACOX1、Ki67","最重要的局部靶细胞之一"],
["Goblet cells","mucin secretion and mucus barrier","mucin不足、microbial encroachment","MUC2/MUC3、secretory differentiation","MUC2、PAS/Alcian blue、goblet-cell number","增强mucus defense"],
["Paneth cells","antimicrobial peptides and innate defense","AMP下降、dysbiosis和translocation","defensins、lysozyme、granule program","LYZ1、DEFA5/6、granule histology","恢复antimicrobial barrier"],
["Macrophages","resolution、efferocytosis、M2-like programs","M1-skewed inflammation","Arg1/MRC1/IL-10、NF-κB repression","CD206、Arg1、TNF、IL-10","促炎症消退和组织修复"],
["Dendritic cells","tolerance、limited antigen presentation","pro-inflammatory DC activation","IL-10/PD-L1、reduced MHC-II/co-stimulation","CD103、PD-L1、CD80/86、IL-12/23","支持tolerogenic DC与Treg"],
["Treg / Th17 axis","immune tolerance","Th17偏移和IL-17 inflammation","FOXP3/IL-10 versus RORγt/IL-17","Treg/Th17 ratio","恢复adaptive immune balance"],
["Neutrophils","controlled recruitment and apoptosis","excess recruitment/ROS/NET injury","chemokine restraint、pro-resolving programs","MPO、Ly6G、CXCL1/2、NET markers","限制急性组织损伤"],
["Fibroblasts / myofibroblasts","stromal quiescence and repair","myofibroblast activation / ECM deposition","TGF-β/SMAD restraint、ECM turnover","α-SMA、COL1A1、FN1、p-SMAD2/3","抗纤维化和结构保护"]
])

h(doc,"Table 4 | Outstanding questions and future directions",2)
table(doc,["主题","当前缺口","为什么重要","建议方法/标志物","优先下一步"],[
["Cell-context specificity","不同上皮、免疫、基质细胞中的PPARγ功能边界不清","解释看似矛盾的疾病效应并指导定向治疗","single-cell/spatial multi-omics；conditional Pparg models；organoids","建立health→inflammation→repair→fibrosis的cell-resolved PPARG atlas"],
["Microbiota–metabolite causality","specific taxa/metabolite→PPARγ的体内因果链仍不足","决定microbiota/postbiotic治疗是否可重复","gnotobiotic models、FMT/defined consortia、stable-isotope tracing","验证明确microbe–metabolite–receptor–phenotype链"],
["Epithelial vs immune PPARγ","两类细胞的相对贡献和互作未定量","决定优先递送到何种细胞","cell-specific knockout、co-culture、spatial readouts","拆分barrier-first与immune-first机制"],
["CRC paradox","PPARγ保护性与潜在促肿瘤作用的切换条件不清","直接影响药物安全和适应证","Apc/p53 models、organoids、stage-resolved human datasets","定义stage/cell/genotype-specific benefit–risk window"],
["Biomarkers","缺乏可预测响应的非侵入性marker panel","临床试验需要机制型分层","metabolites、transcriptome/epigenome、fecal/plasma markers","建立PPARγ activity + barrier + microbiome复合标志物"],
["Intestine-targeted delivery","系统给药副作用限制应用","局部暴露可扩大治疗窗","colon-targeted prodrugs、nanoparticles、pH/microbiota-responsive systems","开发低系统暴露的intestinal PPARγ modulator"],
["Combination therapy","合理组合尚未系统定义","有望减少单药剂量并提高持久缓解","organoid/ex vivo testing、network modeling、adaptive trials","PPARγ modulation + biologics / anti-fibrotics / microbiome strategies"],
["Long-term safety","代谢、骨、心血管及肿瘤背景风险需长期评估","慢性肠病治疗必须长期安全","longitudinal cohorts、real-world evidence、safety biomarker panels","建立长期安全性和context-specific risk framework"]
])

h(doc,"7. 图表与正文的推荐对应关系",1)
bullets(doc,[
    "Figure 1：放在Introduction末尾，承担全文graphical framework。",
    "Figure 2：对应分子机制章节；不要在正文再次大段重复图中文字。",
    "Figure 3：对应上皮代谢/屏障章节，是本文与普通PPARG综述拉开差距的关键图。",
    "Figure 4：对应microbiota–metabolite章节，必须在正文明确association与causality层级。",
    "Figure 5：对应immune/stromal章节，用于把PPARγ从‘anti-inflammatory receptor’提升为cell-state regulator。",
    "Figure 6：对应治疗章节，强调precision modulation和local delivery，而不是笼统建议全身激动PPARγ。",
    "Table 1与疾病章节配套；Table 2与配体/药理章节配套；Table 3承担cell-specific evidence synthesis；Table 4作为Perspective核心表。"
])

h(doc,"8. 投稿前证据与表述控制",1)
bullets(doc,[
    "SCFAs、indoles、bile acids等微生物代谢物可能通过多个宿主受体起效；除非有receptor-specific perturbation，不应全部归因于PPARγ。",
    "IBD中的PPARγ降低、CRC中的PPARγ作用和fibrosis中的药理学效应应分别按human tissue、animal model和cell model分层表述。",
    "5-ASA与PPARγ之间可作为经典转化连接，但不要把所有5-ASA疗效简化为单一PPARγ机制。",
    "TZDs已有成熟代谢病药理学，但其在肠道疾病中的系统性副作用和适应证仍需谨慎讨论。",
    "图中的方向性箭头在正式投稿前应逐条对应claim–evidence ledger；不能把图形美学变成未经验证的因果关系。",
    "若目标是Nature Reviews/Cell风格，正文应该围绕state transition和mechanistic synthesis写作，而不是按研究年份堆叠文献。"
])

h(doc,"9. 一句话创新定位",1)
p(doc,"本综述的核心贡献不是再次证明PPARγ“抗炎”，而是提出一个可检验的肠道生态系统模型：PPARγ通过协调上皮脂肪酸氧化与氧环境、屏障结构、微生物代谢物输入、免疫消退和基质修复，决定肠道从稳态向慢性炎症、纤维化、肿瘤或肠–肝代谢失衡转变；因此未来治疗应从全身强激动转向细胞/组织/疾病阶段特异的精准调节。",10.8,True)

doc.core_properties.title="PPARG通路与肠道相关疾病综述框架与图表整合版"
doc.core_properties.subject="PPARG gut review framework with figures and tables"
doc.core_properties.author="ChatGPT"
doc.save(OUT)

# Validate package
check=Document(OUT)
assert len(check.inline_shapes) == 6, len(check.inline_shapes)
assert len(check.tables) == 4, len(check.tables)
assert OUT.stat().st_size > 500000
print(f"Created: {OUT}")
print(f"Figures: {len(check.inline_shapes)}, Tables: {len(check.tables)}, Paragraphs: {len(check.paragraphs)}, Size: {OUT.stat().st_size}")
