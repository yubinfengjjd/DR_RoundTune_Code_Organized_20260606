from pathlib import Path
import zipfile, shutil, os, sys
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path("chatgpt_exports/cGAS-STING_Autoimmune_Review_Final_Illustrated_CN.docx")
OUT = Path("chatgpt_exports/cGAS-STING_Autoimmune_Review_HighImpact_Benchmark_Final_CN.docx")
TMP = Path("chatgpt_exports/_cgas_build")
TMP.mkdir(parents=True, exist_ok=True)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), "true")
    trPr.append(tblHeader)

def set_font(run, size=10.5, bold=False, color=None, east="宋体", latin="Times New Roman"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)

def add_p(doc, text="", size=10.5, bold=False, align=None, space_after=5, indent=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    return p

def add_bullets(doc, items, level=0):
    for x in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.55 + level*0.55)
        p.paragraph_format.first_line_indent = Cm(-0.3)
        p.paragraph_format.space_after = Pt(2)
        r=p.add_run("• "+x)
        set_font(r, 10.2)

def add_heading(doc, text, level=1):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10 if level==1 else 6)
    p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text)
    if level==1:
        set_font(r, 15, True, "17365D", "黑体")
    elif level==2:
        set_font(r, 12.5, True, "245B78", "黑体")
    else:
        set_font(r, 11, True, "3B6A57", "黑体")
    return p

def make_table(doc, headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers))
    t.style="Table Grid"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=t.rows[0]
    set_repeat_table_header(hdr)
    for i,h in enumerate(headers):
        c=hdr.cells[i]
        set_cell_shading(c, "DCEAF5")
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(h); set_font(r,9.2,True,"17365D","黑体")
    for row in rows:
        cells=t.add_row().cells
        for j,val in enumerate(row):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[j].paragraphs[0]
            p.paragraph_format.space_after=Pt(0)
            r=p.add_run(str(val)); set_font(r,8.8)
    doc.add_paragraph()
    return t

def add_caption(doc, title, legend):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(2)
    r=p.add_run(title); set_font(r,9.8,True,"17365D")
    p2=doc.add_paragraph()
    p2.paragraph_format.space_after=Pt(8)
    r2=p2.add_run(legend); set_font(r2,9.2)

# Extract existing six SVG figures and convert to PNG
fig_dir = TMP/"figs"
fig_dir.mkdir(exist_ok=True)
with zipfile.ZipFile(BASE, "r") as z:
    media=[n for n in z.namelist() if n.startswith("word/media/") and n.lower().endswith(".svg")]
    media=sorted(media)
    for n in media:
        (fig_dir/Path(n).name).write_bytes(z.read(n))

try:
    import cairosvg
except Exception:
    raise SystemExit("cairosvg is required")

pngs=[]
for svg in sorted(fig_dir.glob("*.svg"))[:6]:
    png=svg.with_suffix(".png")
    cairosvg.svg2png(url=str(svg), write_to=str(png), output_width=2400)
    pngs.append(png)

doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)

# default font
style=doc.styles["Normal"]
style.font.name="Times New Roman"
style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
style.font.size=Pt(10.5)

# Title
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("cGAS–STING通路与自身免疫性疾病"); set_font(r,20,True,"17365D","黑体")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("从self-DNA稳态失衡到疾病内表型与精准靶向"); set_font(r,14,True,"245B78","黑体")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("高影响力综述对标版：框架、图表体系与投稿级设计"); set_font(r,11,False,"5B6D7A")
add_p(doc,"建议英文题目：From self-DNA homeostasis failure to autoimmune endotypes: context-dependent cGAS–STING signaling and precision targeting",10.5,False,WD_ALIGN_PARAGRAPH.CENTER,8,False)
add_p(doc,"定位：影响因子 >10；写作与图表设计对标 Nature Reviews Immunology / Nature Reviews Molecular Cell Biology / Cell 等高影响力综述。证据与转化信息更新至 2026 年 9 月。",9.5,False,WD_ALIGN_PARAGRAPH.CENTER,8,False)

doc.add_page_break()

add_heading(doc,"1. 核心定位：为什么不能再写成“通路 + 疾病罗列”",1)
add_p(doc,"现有高影响力cGAS–STING综述已经系统覆盖了经典信号传导、结构调控、通路交叉以及广谱炎症疾病。2025年Nature Reviews Immunology以“通路调控—crosstalk—多种细胞输出”为主线；2021年Nature Reviews Immunology则采用“通路总览—效应机制—无菌炎症—药物机制”的高度压缩叙事；2020年Nature Reviews Molecular Cell Biology把cGAS定位、cGAS激活、STING激活和效应输出分层展开。2026年的Cell综述进一步强调cGAS–STING的情境依赖性和医学转化。")
add_p(doc,"因此，本综述的竞争力不应来自再次总结“cGAS识别DNA并激活STING”，而应建立一个自身免疫领域专属的解释框架：self-DNA生成与暴露增加、清除与空间隔离失败、cGAMP/STING信号终止不足、跨细胞传播以及组织特异性免疫回路共同决定疾病内表型。")
add_bullets(doc,[
    "主问题1：什么样的self-DNA真正具有免疫刺激性，来源于哪里，又为什么没有被及时清除？",
    "主问题2：哪些检查点决定cGAS–STING是短暂宿主防御，还是持续病理性炎症？",
    "主问题3：信号如何从一个细胞扩展为B/T细胞、髓系细胞、内皮细胞和成纤维细胞共同参与的组织网络？",
    "主问题4：不同疾病是否可被重构为少数机制内表型，而不是简单按病名罗列？",
    "主问题5：哪些患者、哪些组织、哪个疾病阶段真正适合直接cGAS/STING抑制？"
])

add_heading(doc,"2. 高分综述反向工程：可直接借鉴的结构与图表语言",1)
benchmark_rows=[
    ["Nature Reviews Immunology 2025","6幅主图","Overview → regulation → innate crosstalk → metabolism → ageing → tumour–immune interactions","不是按疾病罗列，而是围绕“调控与交叉网络”逐层加深。"],
    ["Nature Reviews Immunology 2021","4幅主图","Pathway → effector/intercellular cGAMP → sterile inflammation → inhibitors","图少但每张图承担一个完整逻辑模块，避免装饰性图。"],
    ["Nature Reviews Molecular Cell Biology 2020","5幅主图","Canonical pathway → localization/ligands → cGAS activation → STING activation → effectors","把同一通路拆成空间、结构和效应层级，机制深度高。"],
    ["Cell 2026","情境型综述","Mechanism → physiological functions → disease implications → context-dependent therapy","强调保护性与病理性作用并存，治疗必须匹配情境。"]
]
make_table(doc,["对标综述","图表规模","叙事骨架","对本综述的启示"],benchmark_rows)

add_heading(doc,"3. 本综述的独特主线",1)
add_p(doc,"建议全文用一条统一逻辑串联：")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("self-DNA dysregulation  →  checkpoint failure  →  persistent cGAS–STING activation  →  intercellular propagation  →  autoimmune endotypes  →  precision targeting")
set_font(r,12,True,"8A3150")
add_bullets(doc,[
    "把SAVI、Aicardi–Goutières syndrome等单基因干扰素病作为“机制因果锚点”，但明确它们属于interferonopathy/autoinflammation，而不是把它们与常规自身免疫病混为一类。",
    "系统性红斑狼疮作为复杂自身免疫中的“原型疾病”，用于展示self-DNA、NETs、免疫复合物、I型干扰素与B细胞自身抗体之间的闭环。",
    "Sjögren’s syndrome、dermatomyositis、systemic sclerosis、rheumatoid arthritis、vasculitis及CNS autoimmune inflammation用于体现组织和细胞情境依赖性。",
    "把“IFN signature ≠ 直接证明cGAS–STING依赖”写进全文方法论，防止综述过度推断。"
])

add_heading(doc,"4. 推荐正文框架",1)
outline=[
("4.1 Introduction: from antiviral defense to self-DNA-driven autoimmunity",[
"提出核心悖论：同一条DNA感知通路如何从保护性宿主防御转为慢性自身免疫。",
"指出该领域已有大量广谱综述，本综述聚焦self-DNA稳态、检查点、组织传播和疾病内表型。"]),
("4.2 Molecular grammar of cGAS–STING signaling",[
"cGAS识别dsDNA、2′3′-cGAMP生成、STING从ER向ERGIC/Golgi转运、TBK1–IRF3与IKK–NF-κB。",
"补充nuclear cGAS、相分离、STING trafficking、非经典自噬/细胞死亡/炎症小体等输出。"]),
("4.3 Sources and fates of immunostimulatory self-DNA",[
"核DNA损伤、微核、染色体不稳定和RNA:DNA hybrids。",
"线粒体应激与mtDNA释放；NET-derived DNA；凋亡/坏死残骸；潜在retroelement来源。",
"强调DNA来源、长度、持续时间、细胞定位和蛋白复合状态共同决定免疫原性。"]),
("4.4 Homeostatic checkpoints that restrain the pathway",[
"核内/染色质隔离与DNA repair；PINK1–Parkin等线粒体质量控制。",
"TREX1、RNase H2、DNase II、SAMHD1及autophagy–lysosome清除。",
"cGAMP运输与降解（ENPP1/ENPP3等）；STIM1、COPA、AP-1/ESCRT等STING转运与终止机制。"]),
("4.5 From cell-intrinsic sensing to tissue-level autoimmune circuits",[
"树突状细胞、单核/巨噬细胞、中性粒细胞、B/T细胞、内皮细胞、成纤维细胞的差异化角色。",
"cGAMP、I型干扰素、趋化因子及细胞死亡导致的跨细胞传播与feed-forward loops。",
"重点讨论何时信号仍是局部防御，何时跨过阈值转为持续自身免疫。"]),
("4.6 Mechanistic endotypes across autoimmune diseases",[
"DNA-clearance / interferon endotype：SLE为原型。",
"Vascular–endothelial endotype：SAVI作为因果锚点，并讨论vasculitis/systemic sclerosis。",
"Epithelial/glandular endotype：Sjögren’s syndrome。",
"Muscle/skin injury endotype：dermatomyositis。",
"Synovial/mesenchymal endotype：rheumatoid arthritis。",
"CNS/glial endotype：神经自身免疫/神经炎症。"]),
("4.7 Crosstalk and non-canonical outputs",[
"TLR7/9–MyD88、NLRP3 inflammasome、JAK–STAT、DNA damage response。",
"mitochondrial ROS与代谢重编程、autophagy–lysosome、senescence/SASP。",
"区分IFN-dependent与IFN-independent STING输出。"]),
("4.8 Therapeutic targeting and the pathogenic-versus-protective window",[
"减少self-DNA/NET负荷；增强DNA清除；直接cGAS抑制；cGAMP传输/降解调控；STING拮抗/降解；下游IFN/JAK/TBK1干预。",
"组织靶向递送、组合治疗、安全性监测以及感染/抗肿瘤免疫风险。",
"以VENT-03进入CLE/SLE Phase 2a为直接cGAS临床转化的重要里程碑。"]),
("4.9 Biomarkers, trial design and outstanding questions",[
"基因、IFN signature、cGAMP、p-STING/p-TBK1、cfDNA/mtDNA、NET markers及组织空间信号。",
"未来临床试验从diagnosis-based enrollment转向mechanism-based enrollment。",
"回答which patient, which tissue, which stage, which node四个问题。"])
]
for title, bullets in outline:
    add_heading(doc,title,2); add_bullets(doc,bullets)

doc.add_page_break()
add_heading(doc,"5. 机制图体系（6幅主图）",1)

fig_titles=[
"Figure 1 | From self-DNA homeostasis failure to cGAS–STING-driven autoimmune pathology",
"Figure 2 | Molecular checkpoints that preserve self-DNA tolerance and restrain cGAS–STING signaling",
"Figure 3 | From cell-intrinsic DNA sensing to multicellular autoimmune circuits",
"Figure 4 | Disease-context map of cGAS–STING across autoimmune and interferon-mediated disorders",
"Figure 5 | Crosstalk and non-canonical outputs of cGAS–STING in autoimmunity",
"Figure 6 | Therapeutic targeting of the cGAS–STING axis: intervention points and translational roadmap"
]
fig_legends=[
"以self-DNA来源、核心cGAS–cGAMP–STING级联、主要效应细胞和疾病谱为一体化总览。图的目的不是再画一次教科书通路，而是把“DNA稳态失败—持续信号—组织损伤”作为全文第一视觉主线。",
"展示核/染色质完整性、线粒体质量控制、核酸酶、autophagy–lysosome、cGAMP handling以及STING trafficking/termination等多级制动。该图对应全文最重要的创新概念：自身免疫并非只因通路过强，也可能因多个homeostatic brakes同时失效。",
"将树突状细胞、巨噬细胞、中性粒细胞、内皮/基质细胞、B细胞和T细胞连接为一个组织级feed-forward circuit，突出免疫复合物、NETs、I型干扰素和细胞死亡产生的新一轮self-DNA。",
"比较SAVI/干扰素病、SLE、Sjögren’s syndrome、dermatomyositis、systemic sclerosis、RA、vasculitis与神经自身免疫。图中应明确证据强度并不相同，SAVI/AGS用于机制锚定，而非等同于常规自身免疫病。",
"以TLR7/9、NLRP3、JAK–STAT、mitochondrial ROS/metabolism、autophagy–lysosome和DNA damage–senescence为六个环绕模块，体现cGAS–STING是网络枢纽而不是孤立线性通路。",
"从上游self-DNA/NET负荷、DNA清除、cGAS、cGAMP、STING到TBK1/IFN/JAK分层展示治疗节点，并在下方增加biomarker selection、tissue targeting、combination therapy与safety monitoring，形成机制到临床的闭环。"
]
for i,png in enumerate(pngs):
    doc.add_picture(str(png), width=Inches(6.7))
    doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc,fig_titles[i],fig_legends[i])

doc.add_page_break()
add_heading(doc,"6. 表格体系（4个原生可编辑Word表格）",1)

# Table 1
add_heading(doc,"Table 1 | 高影响力综述的设计范式与本综述的差异化策略",2)
make_table(doc,["模块","高分综述常见处理","本综述建议升级"],[
["通路基础","经典cGAS→cGAMP→STING→TBK1/IRF3","压缩基础知识，把篇幅让给self-DNA稳态与检查点"],
["图表","少而重、每图解决一个概念","6幅主图分别承担总览、制动、传播、疾病、crosstalk、治疗"],
["疾病章节","按疾病逐一讨论","改为mechanistic endotypes，再在每个endotype中放代表疾病"],
["证据表达","机制与疾病并列","增加human genetics / patient tissue / perturbation / association四级证据语言"],
["治疗","列药物与靶点","增加patient–tissue–stage–biomarker四维治疗窗口"]
])

# Table 2
add_heading(doc,"Table 2 | cGAS–STING相关自身免疫疾病的机制证据矩阵",2)
rows=[
["SLE","核/mtDNA、NETs、清除不足","pDC/髓系/B细胞/肾皮肤","IFN-I、免疫复合物、CXCL10","强机制证据但高度异质","证明cGAS依赖而非仅IFN-high"],
["Sjögren’s syndrome","腺上皮DNA应激/mtDNA","腺上皮、pDC、B细胞","局部IFN程序、腺体炎症","中等/新兴","组织内因果与亚群分层"],
["Dermatomyositis","肌细胞损伤、mtDNA","肌细胞、皮肤、髓系","IFN-rich inflammation","新兴","疾病亚型与cGAS依赖性"],
["Systemic sclerosis","细胞损伤、衰老、mtDNA","内皮、成纤维细胞","血管炎症、纤维化","新兴","保护性/促纤维化作用拆分"],
["Rheumatoid arthritis","NETs、mtDNA、滑膜损伤","巨噬、成纤维样滑膜细胞","炎症、组织破坏","情境依赖","确定STING-high滑膜亚群"],
["Vasculitis","NETs、内皮损伤、DAMPs","中性粒、内皮、单核细胞","血管炎症、IFN","新兴","不同血管炎亚型差异"],
["Autoimmune neuroinflammation","神经/胶质损伤DNA、mtDNA","microglia、astrocyte","局部IFN、神经炎症","新兴","CNS内源性与外周信号区分"],
["SAVI / AGS","STING GOF或DNA代谢基因缺陷","内皮/髓系等","持续IFN与全身炎症","人类遗传因果锚点","应定义为interferonopathy/autoinflammation，不与常规自身免疫等同"]
]
make_table(doc,["疾病/锚点","主要DNA来源/缺陷","主要细胞/组织","关键输出","证据成熟度","最重要缺口"],rows)

# Table 3
add_heading(doc,"Table 3 | 自身DNA稳态检查点、可测标志物与验证策略",2)
rows=[
["核/染色质隔离","核膜、nucleosome tethering、BAF、DNA repair","micronuclei、cGAS localization、γH2AX","成像+DNA damage perturbation"],
["线粒体质量控制","PINK1–Parkin/mitophagy、ROS控制","cytosolic mtDNA、oxidized mtDNA、mitochondrial stress","mitophagy manipulation+mtDNA depletion/rescue"],
["核酸清除","TREX1、RNase H2、DNase II、SAMHD1","cfDNA/cytosolic DNA、nuclease activity","genetic rescue / nuclease replacement"],
["cGAMP处理","ENPP1/ENPP3、transporters/gap junctions","cGAMP水平、transport activity","transport/degradation perturbation"],
["STING ER保留/转运","STIM1、COPA、ER–Golgi traffic","STING localization、p-STING/p-TBK1","trafficking perturbation"],
["信号终止","AP-1/ESCRT、autophagy–lysosome、phosphatases/DUBs","STING turnover、lysosomal localization","degradation blockade/rescue"],
["下游放大","IFNAR–JAK–STAT、NF-κB、NLRP3","IFN gene score、CXCL10、IL-1β/IL-18","pathway-specific inhibition"]
]
make_table(doc,["层级","代表性检查点","候选标志物","更强的因果验证"],rows)

# Table 4
add_heading(doc,"Table 4 | 2026年转化治疗图谱：靶点、代表策略与证据边界",2)
rows=[
["减少self-DNA/NETs","PAD4/NETosis抑制、DNase/清除增强","多数为前临床/转化研究","从源头减配体，但存在通路冗余和感染风险"],
["cGAS","VENT-03（口服cGAS inhibitor）","CLE±SLE Phase 2a，NCT07260877，招募中","目前最重要的直接cGAS临床验证；主要终点之一为皮肤IFN gene signature"],
["cGAS","RU.521、PF-06928215等研究工具/先导化合物","前临床","物种差异、细胞内暴露与生化效力差距需谨慎"],
["STING","H-151、C-176/C-178等","前临床","直接拮抗逻辑清晰，但尚缺成熟自身免疫临床验证"],
["STING degradation/trafficking","PROTAC/lysosome、ER–Golgi trafficking干预","早期研究","可能延长抑制并提高选择性，需证明组织与长期安全性"],
["IFNAR","anifrolumab","SLE已获批（下游验证）","验证IFN轴可药物化，但并非STING特异性"],
["JAK","baricitinib/tofacitinib等","多种自身免疫病已临床应用","可降低IFN相关输出，但属于广谱下游调节"]
]
make_table(doc,["靶点层级","代表策略/药物","截至2026-09状态","解读与限制"],rows)

add_heading(doc,"7. 投稿级“防审稿人攻击”规则",1)
add_bullets(doc,[
    "不要用IFN signature直接等同于cGAS–STING激活；至少区分TLR7/9、MDA5等其他IFN来源。",
    "不要把所有STING信号都写成cGAS依赖；应保留cGAS-independent STING activation和非经典输出的可能性。",
    "SAVI与AGS用于证明DNA/STING轴具有因果能力，但应明确其分类与常见自身免疫病不同。",
    "疾病章节中分开写patient evidence、animal model evidence和in vitro evidence，避免跨层级外推。",
    "直接cGAS/STING抑制与IFN/JAK下游药物必须分开，后者只能作为pathway druggability的临床支持。",
    "治疗部分必须讨论感染、抗病毒、抗肿瘤监视和组织修复受损风险。",
    "每幅图下方可加一句‘evidence strength varies by disease and cell context’，避免图形造成过度因果暗示。"
])

add_heading(doc,"8. 建议的最终文章标题与卖点",1)
add_bullets(doc,[
    "首选：From self-DNA homeostasis failure to autoimmune endotypes: context-dependent cGAS–STING signaling and precision targeting",
    "备选：The cGAS–STING axis in autoimmune disease: checkpoints, multicellular propagation and therapeutic windows",
    "中文：cGAS–STING通路与自身免疫性疾病：从自身DNA稳态失衡到疾病内表型与精准靶向"
])
add_p(doc,"一句话卖点：这不是另一篇“cGAS–STING与疾病”的综述，而是用self-DNA homeostasis、checkpoint failure、intercellular propagation和mechanistic endotypes解释为什么同一通路在不同自身免疫病中产生不同病理输出，并据此提出可检验的精准治疗框架。",10.8,True)

add_heading(doc,"9. 对标文献（公开资料核对，更新至2026-09）",1)
refs=[
"Zhang Z, Zhang C. Regulation of cGAS–STING signalling and its diversity of cellular outcomes. Nature Reviews Immunology. 2025;25:425–444.",
"Decout A, Katz JD, Venkatraman S, Ablasser A. The cGAS–STING pathway as a therapeutic target in inflammatory diseases. Nature Reviews Immunology. 2021;21:548–569.",
"Hopfner KP, Hornung V. Molecular mechanisms and cellular functions of cGAS–STING signalling. Nature Reviews Molecular Cell Biology. 2020;21:501–521.",
"Motwani M, Pesiridis S, Fitzgerald KA. DNA sensing by the cGAS–STING pathway in health and disease. Nature Reviews Genetics. 2019;20:657–674.",
"Skopelja-Gardner S, An J, Elkon KB. Role of the cGAS–STING pathway in systemic and organ-specific diseases. Nature Reviews Nephrology. 2022;18:558–572.",
"Zhang B, Xu P, Ablasser A. Regulation of the cGAS-STING Pathway. Annual Review of Immunology. 2025;43:667–692.",
"The cGAS-STING pathway: Mechanism and medical implications. Cell. 2026;189:3849–3870.",
"ClinicalTrials.gov NCT07260877 (AERIS): VENT-03 in active cutaneous lupus erythematosus with or without SLE; Phase 2a, recruiting as of 2026."
]
for x in refs:
    add_p(doc,x,9.2,False,None,2,False)

doc.core_properties.title="cGAS–STING通路与自身免疫性疾病：高影响力综述对标版"
doc.core_properties.subject="Review framework, figures and tables"
doc.core_properties.author="ChatGPT"
doc.save(OUT)

# verification
check=Document(OUT)
assert len(check.paragraphs)>80
assert len(check.tables)>=5
assert OUT.stat().st_size>200000
print(OUT)
print("paragraphs",len(check.paragraphs),"tables",len(check.tables),"size",OUT.stat().st_size)
