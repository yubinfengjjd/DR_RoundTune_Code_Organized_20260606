from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_AUTO_SIZE
from pathlib import Path
import math
import requests

OUT = Path("chatgpt_exports/药物小侦探_一年级科普课_30分钟.pptx")
OUT.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
NAVY = RGBColor(28, 62, 95)
BLUE = RGBColor(79, 149, 230)
SKY = RGBColor(225, 242, 255)
CORAL = RGBColor(241, 111, 103)
PINK = RGBColor(255, 225, 230)
GREEN = RGBColor(71, 177, 122)
MINT = RGBColor(224, 247, 236)
YELLOW = RGBColor(252, 201, 73)
CREAM = RGBColor(255, 249, 235)
PURPLE = RGBColor(139, 111, 215)
LAV = RGBColor(239, 233, 255)
GRAY = RGBColor(101, 116, 130)
DARK = RGBColor(35, 45, 58)
WHITE = RGBColor(255,255,255)
RED = RGBColor(220, 64, 64)

def add_bg(slide, color=WHITE):
    sh=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height)
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background()
    slide.shapes._spTree.remove(sh._element); slide.shapes._spTree.insert(2, sh._element)

def add_text(slide, text, x,y,w,h, size=26, color=DARK, bold=False, align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear(); tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; r.font.name=font
    tf.word_wrap=True
    return box

def title(slide, text, subtitle=None, color=NAVY):
    add_text(slide,text,0.6,0.35,12.1,0.65,32,color,True,PP_ALIGN.CENTER)
    if subtitle:
        add_text(slide,subtitle,1.1,1.02,11.1,0.45,15,GRAY,False,PP_ALIGN.CENTER)

def rounded(slide,x,y,w,h,fill,line=None,radius=None):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line; s.line.width=Pt(1.2)
    return s

def circle(slide,x,y,d,fill,line=None):
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb=fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb=line
    return s

def pill(slide,x,y,w=1.7,h=0.7,left=BLUE,right=CORAL,face=True):
    # one rounded pill, split with a center white line
    base=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    base.fill.solid(); base.fill.fore_color.rgb=left; base.line.color.rgb=WHITE; base.line.width=Pt(1.2)
    r=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x+w/2),Inches(y),Inches(w/2),Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb=right; r.line.fill.background()
    # cover center left-rounded part of right with rect
    rr=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+w/2),Inches(y),Inches(w/4),Inches(h))
    rr.fill.solid(); rr.fill.fore_color.rgb=right; rr.line.fill.background()
    if face:
        circle(slide,x+0.45,y+0.2,0.08,DARK)
        circle(slide,x+0.7,y+0.2,0.08,DARK)
        smile=slide.shapes.add_shape(MSO_SHAPE.ARC,Inches(x+0.48),Inches(y+0.28),Inches(0.3),Inches(0.18))
        smile.line.color.rgb=DARK; smile.line.width=Pt(1.5); smile.fill.background()
    return base

def add_cloud(slide,x,y,w,h,fill=SKY):
    for dx,dy,d in [(0.1,0.25,0.7),(0.55,0.05,0.95),(1.2,0.2,0.7)]:
        circle(slide,x+dx,y+dy,d,fill)
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x+0.15),Inches(y+0.45),Inches(w-0.3),Inches(h-0.45))
    s.fill.solid(); s.fill.fore_color.rgb=fill; s.line.fill.background()

def arrow(slide,x1,y1,x2,y2,color=BLUE,width=3):
    line=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x1),Inches(y1),Inches(x2-x1),Inches(max(0.25,y2-y1 if y2>y1 else 0.25)))
    line.fill.solid(); line.fill.fore_color.rgb=color; line.line.fill.background()
    return line

def add_interaction(slide, text):
    rounded(slide,0.65,6.73,12.0,0.5,CREAM,YELLOW)
    add_text(slide,"🎤 互动："+text,0.9,6.78,11.5,0.38,16,NAVY,True)

def add_footer(slide, n, mins):
    add_text(slide,f"{n:02d}",12.45,0.15,0.45,0.35,12,GRAY,True,PP_ALIGN.RIGHT)
    add_text(slide,f"约 {mins} min",11.45,0.15,0.85,0.35,11,GRAY,False,PP_ALIGN.RIGHT)

def add_character(slide,x,y):
    pill(slide,x,y,2.1,0.85,BLUE,CORAL,True)
    # goggles
    circle(slide,x+0.32,y+0.08,0.28,WHITE,NAVY); circle(slide,x+0.66,y+0.08,0.28,WHITE,NAVY)
    # legs
    l1=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+0.45),Inches(y+0.78),Inches(0.08),Inches(0.35)); l1.fill.solid(); l1.fill.fore_color.rgb=NAVY; l1.line.fill.background()
    l2=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+1.5),Inches(y+0.78),Inches(0.08),Inches(0.35)); l2.fill.solid(); l2.fill.fore_color.rgb=NAVY; l2.line.fill.background()

def safe_bullet(slide, text, x,y,color=GREEN):
    circle(slide,x,y,0.28,color)
    add_text(slide,"✓",x+0.03,y-0.01,0.23,0.23,16,WHITE,True,PP_ALIGN.CENTER)
    add_text(slide,text,x+0.4,y-0.05,5.5,0.45,20,DARK,False)

# Slide 1
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY)
add_text(s,"药物小侦探",0.6,0.9,12.1,0.8,44,NAVY,True,PP_ALIGN.CENTER)
add_text(s,"一颗小小的药，是怎样帮助身体的？",1.2,1.75,10.9,0.55,24,CORAL,True,PP_ALIGN.CENTER)
add_character(s,5.55,3.0)
for x,c in [(2.0,GREEN),(3.25,YELLOW),(8.6,PURPLE),(10.0,CORAL)]:
    circle(s,x,3.25,0.7,c)
    add_text(s,"?",x+0.12,3.26,0.45,0.4,28,WHITE,True,PP_ALIGN.CENTER)
rounded(s,3.1,4.65,7.1,1.05,WHITE,BLUE)
add_text(s,"今天我们一起破案：药物为什么能帮助身体？",3.35,4.82,6.6,0.65,22,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,1,1)

# Slide 2
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"你见过哪些“药物小伙伴”？")
items=[("药片",BLUE),("胶囊",CORAL),("药水",GREEN),("喷雾",PURPLE),("药膏",YELLOW),("贴剂",RGBColor(76,190,200))]
for i,(name,c) in enumerate(items):
    x=0.8+(i%3)*4.1; y=1.7+(i//3)*2.15
    rounded(s,x,y,3.5,1.55,RGBColor(min(c[0]+165,255),min(c[1]+165,255),min(c[2]+165,255)),c)
    if name=="胶囊": pill(s,x+0.55,y+0.42,1.35,0.55,BLUE,CORAL,False)
    elif name=="药片": circle(s,x+0.65,y+0.4,0.65,c)
    elif name=="药水":
        b=slide=None
        rect=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x+0.7),Inches(y+0.35),Inches(0.65),Inches(0.85)); rect.fill.solid(); rect.fill.fore_color.rgb=c; rect.line.fill.background()
        cap=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+0.82),Inches(y+0.2),Inches(0.4),Inches(0.18)); cap.fill.solid(); cap.fill.fore_color.rgb=NAVY; cap.line.fill.background()
    elif name=="喷雾":
        rect=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+0.7),Inches(y+0.35),Inches(0.55),Inches(0.8)); rect.fill.solid(); rect.fill.fore_color.rgb=c; rect.line.fill.background()
        add_text(s,"💨",x+1.3,y+0.35,0.8,0.6,28,NAVY)
    elif name=="药膏":
        tri=s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x+0.55),Inches(y+0.35),Inches(0.9),Inches(0.8)); tri.rotation=90; tri.fill.solid(); tri.fill.fore_color.rgb=c; tri.line.fill.background()
    else:
        rounded(s,x+0.55,y+0.45,1.1,0.55,c,c)
    add_text(s,name,x+1.95,y+0.48,1.2,0.5,22,NAVY,True,PP_ALIGN.CENTER)
add_interaction(s,"举手说一说：你见过哪一种？哪一种不用吞下去？"); add_footer(s,2,2)

# Slide 3 body city
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM); title(s,"我们的身体，是一座超级城市！")
# torso
rounded(s,4.3,1.5,4.7,4.7,RGBColor(255,238,225),CORAL)
circle(s,6.0,0.95,1.2,RGBColor(255,224,205),CORAL)
# heart
circle(s,5.1,2.2,0.7,CORAL); add_text(s,"❤",5.18,2.2,0.55,0.45,28,WHITE,True,PP_ALIGN.CENTER)
# lungs
circle(s,6.6,2.1,0.75,SKY,BLUE); circle(s,7.35,2.1,0.75,SKY,BLUE)
# intestine
rounded(s,5.55,3.65,2.4,1.35,MINT,GREEN); add_text(s,"肠道\\n营养工厂",5.9,3.85,1.7,0.75,18,GREEN,True,PP_ALIGN.CENTER)
# immune police
for dx in [0,0.65,1.3]:
    circle(s,4.7+dx,5.25,0.48,LAV,PURPLE)
add_text(s,"免疫小警察",4.6,5.78,2.2,0.4,18,PURPLE,True,PP_ALIGN.CENTER)
# labels
rounded(s,0.65,1.8,2.8,1.0,PINK,CORAL); add_text(s,"心脏＝泵站",0.9,2.03,2.3,0.45,24,CORAL,True,PP_ALIGN.CENTER)
rounded(s,9.85,1.8,2.8,1.0,SKY,BLUE); add_text(s,"肺＝空气站",10.1,2.03,2.3,0.45,24,BLUE,True,PP_ALIGN.CENTER)
rounded(s,0.65,4.05,2.8,1.0,MINT,GREEN); add_text(s,"肠道＝营养工厂",0.84,4.28,2.4,0.45,21,GREEN,True,PP_ALIGN.CENTER)
rounded(s,9.85,4.05,2.8,1.0,LAV,PURPLE); add_text(s,"免疫＝警察局",10.05,4.28,2.4,0.45,22,PURPLE,True,PP_ALIGN.CENTER)
add_interaction(s,"我指一个器官，你们猜猜它在“身体城市”里做什么工作？"); add_footer(s,3,2)

# Slide 4
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"生病时，身体城市发生了什么？")
# healthy vs sick
rounded(s,0.8,1.5,5.6,4.7,MINT,GREEN)
add_text(s,"健康城市",1.1,1.72,5.0,0.5,28,GREEN,True,PP_ALIGN.CENTER)
for i in range(6):
    circle(s,1.5+i*0.65,2.8+(i%2)*0.4,0.35,BLUE if i%2==0 else GREEN)
add_text(s,"✓ 秩序井然\\n✓ 屏障完整\\n✓ 免疫小警察巡逻",1.4,4.15,4.4,1.35,22,DARK,False)
rounded(s,6.95,1.5,5.6,4.7,PINK,CORAL)
add_text(s,"生病城市",7.25,1.72,5.0,0.5,28,CORAL,True,PP_ALIGN.CENTER)
for i in range(8):
    circle(s,7.65+i*0.5,2.75+(i%3)*0.27,0.32,CORAL if i%2==0 else PURPLE)
add_text(s,"⚠ 病菌可能来捣乱\\n⚠ 身体会发炎\\n⚠ 有些功能会“出小故障”",7.45,4.15,4.5,1.35,21,DARK)
add_interaction(s,"找不同：右边的身体城市，哪里和左边不一样？"); add_footer(s,4,2)

# Slide 5
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY); title(s,"药物不是魔法，是“有任务的小帮手”")
add_character(s,1.1,2.4)
rounded(s,4.0,1.65,7.9,3.95,WHITE,BLUE)
tasks=[("退烧","把太高的体温降下来",CORAL),("止痛","让疼痛信号变弱",PURPLE),("抗感染","帮助对付某些病原体",GREEN)]
for i,(a,b,c) in enumerate(tasks):
    circle(s,4.55,2.05+i*1.05,0.55,c)
    add_text(s,str(i+1),4.7,2.11+i*1.05,0.25,0.25,17,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,a,5.25,1.98+i*1.05,1.2,0.4,23,c,True)
    add_text(s,b,6.55,1.98+i*1.05,4.6,0.45,21,DARK)
add_text(s,"不同药物，有不同任务。",3.2,5.95,6.9,0.5,28,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,5,2)

# Slide 6
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"为什么药物有这么多“外形”？")
forms=[("药片","方便携带",BLUE),("胶囊","保护里面的小药物",CORAL),("药水","有些小朋友更容易喝",GREEN),("喷雾","直接到鼻子/喉咙等地方",PURPLE),("药膏","涂在皮肤上",YELLOW),("贴剂","贴在皮肤上慢慢工作",RGBColor(76,190,200))]
for i,(n,desc,c) in enumerate(forms):
    x=0.7+(i%3)*4.2; y=1.55+(i//3)*2.45
    rounded(s,x,y,3.75,1.85,RGBColor(247,250,253),c)
    circle(s,x+0.35,y+0.45,0.65,c)
    add_text(s,n,x+1.2,y+0.35,2.1,0.45,24,NAVY,True)
    add_text(s,desc,x+1.2,y+0.83,2.2,0.7,17,GRAY)
add_interaction(s,"猜一猜：如果皮肤擦伤，哪一类药更可能“直接到现场”？"); add_footer(s,6,2)

# Slide 7 journey
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM); title(s,"一颗药去旅行：它会经过哪里？")
stops=[("1","嘴巴",BLUE),("2","胃",CORAL),("3","肠道",GREEN),("4","血液",PURPLE),("5","目标地点",YELLOW)]
for i,(num,name,c) in enumerate(stops):
    x=0.7+i*2.45
    circle(s,x,2.8,1.15,c)
    add_text(s,num,x+0.33,2.94,0.5,0.35,26,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,name,x-0.25,4.05,1.65,0.5,22,NAVY,True,PP_ALIGN.CENTER)
    if i<4:
        arr=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+1.3),Inches(3.12),Inches(0.95),Inches(0.35))
        arr.fill.solid(); arr.fill.fore_color.rgb=GRAY; arr.line.fill.background()
add_character(s,5.3,5.1)
add_text(s,"“我要找到我的任务地点！”",7.55,5.25,4.1,0.55,23,CORAL,True)
add_interaction(s,"全班一起用手指沿着路线走一遍：嘴巴 → 胃 → 肠道 → 血液 → 目标"); add_footer(s,7,2)

# Slide 8 lock/key
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"药物怎么找到“目标”？——钥匙和锁")
# keys
for i,c in enumerate([BLUE,GREEN,CORAL]):
    rounded(s,0.9,1.65+i*1.45,3.3,1.05,RGBColor(248,250,252),c)
    circle(s,1.2,1.92+i*1.45,0.45,c)
    stem=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(1.6),Inches(2.08+i*1.45),Inches(1.15),Inches(0.12)); stem.fill.solid(); stem.fill.fore_color.rgb=c; stem.line.fill.background()
    add_text(s,f"钥匙 {i+1}",2.9,1.9+i*1.45,0.9,0.4,19,NAVY,True,PP_ALIGN.CENTER)
# locks
shapes=[MSO_SHAPE.OVAL,MSO_SHAPE.ISOSCELES_TRIANGLE,MSO_SHAPE.DIAMOND]
for i,(shape,c) in enumerate(zip(shapes,[CORAL,BLUE,GREEN])):
    rounded(s,8.5,1.65+i*1.45,3.3,1.05,RGBColor(248,250,252),c)
    lock=s.shapes.add_shape(shape,Inches(8.9),Inches(1.88+i*1.45),Inches(0.55),Inches(0.55)); lock.fill.solid(); lock.fill.fore_color.rgb=c; lock.line.fill.background()
    add_text(s,f"小锁 {i+1}",9.8,1.9+i*1.45,1.3,0.4,19,NAVY,True,PP_ALIGN.CENTER)
add_text(s,"药物常常要找到合适的“靶点”",4.35,2.85,3.75,0.8,27,PURPLE,True,PP_ALIGN.CENTER)
add_interaction(s,"配对游戏：哪把钥匙最可能打开哪把锁？"); add_footer(s,8,2)

# Slide 9 gates
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY); title(s,"身体里还有很多“大门”和“关卡”")
# three gates
gates=[("肠道屏障","不是所有东西都能进去",GREEN),("血液运输","需要搭上“运输车”",BLUE),("特殊保护区","有些地方更难进入",PURPLE)]
for i,(n,d,c) in enumerate(gates):
    x=0.9+i*4.15
    rounded(s,x,1.9,3.6,3.6,WHITE,c)
    # gate icon
    for k in range(3):
        r=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x+0.75+k*0.55),Inches(2.55),Inches(0.18),Inches(1.25))
        r.fill.solid(); r.fill.fore_color.rgb=c; r.line.fill.background()
    add_text(s,n,x+0.35,4.05,2.9,0.45,23,c,True,PP_ALIGN.CENTER)
    add_text(s,d,x+0.45,4.55,2.7,0.55,17,GRAY,False,PP_ALIGN.CENTER)
add_text(s,"科研工作者会想办法让药物“走对路”。",1.5,5.92,10.3,0.55,27,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,9,2)

# Slide 10 drug discovery
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"科学家怎样发明一种新药？")
steps=[("发现问题",CORAL),("找到目标",PURPLE),("设计候选",BLUE),("做实验",GREEN),("检查安全",YELLOW),("帮助病人",RGBColor(76,190,200))]
for i,(name,c) in enumerate(steps):
    x=0.55+i*2.08
    circle(s,x+0.3,2.35,1.0,c)
    add_text(s,str(i+1),x+0.62,2.56,0.35,0.28,22,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,name,x,3.55,1.6,0.55,19,NAVY,True,PP_ALIGN.CENTER)
    if i<5:
        ar=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x+1.4),Inches(2.7),Inches(0.55),Inches(0.28)); ar.fill.solid(); ar.fill.fore_color.rgb=GRAY; ar.line.fill.background()
rounded(s,1.1,4.55,11.1,1.05,CREAM,YELLOW)
add_text(s,"科学不是“一次就成功”，而是：提问 → 实验 → 改进 → 再实验",1.45,4.78,10.4,0.55,23,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,10,2)

# Slide 11 scientist parent
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM); title(s,"科研工作者的一天：我在做什么？")
activities=[("读资料","发现别人已经知道什么",BLUE),("做实验","看看想法对不对",GREEN),("看数据","寻找规律和答案",PURPLE),("改方案","失败也能告诉我们新东西",CORAL)]
for i,(a,b,c) in enumerate(activities):
    y=1.55+i*1.15
    circle(s,0.95,y+0.15,0.65,c)
    add_text(s,str(i+1),1.17,y+0.28,0.25,0.2,18,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,a,1.9,y,1.55,0.45,23,c,True)
    add_text(s,b,3.55,y,5.4,0.48,20,DARK)
# mini lab
rounded(s,9.3,1.65,3.2,4.45,WHITE,BLUE)
add_text(s,"小实验室",9.65,1.95,2.5,0.45,24,NAVY,True,PP_ALIGN.CENTER)
# flasks
for j,c in enumerate([BLUE,GREEN,CORAL]):
    body=s.shapes.add_shape(MSO_SHAPE.TRAPEZOID,Inches(9.75+j*0.75),Inches(3.0),Inches(0.6),Inches(0.8)); body.fill.solid(); body.fill.fore_color.rgb=c; body.line.color.rgb=NAVY
    neck=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(9.95+j*0.75),Inches(2.62),Inches(0.2),Inches(0.4)); neck.fill.solid(); neck.fill.fore_color.rgb=WHITE; neck.line.color.rgb=NAVY
add_text(s,"好奇心 + 认真 + 耐心",9.55,4.7,2.7,0.8,20,PURPLE,True,PP_ALIGN.CENTER)
add_interaction(s,"你觉得科学家最重要的能力是什么？聪明？耐心？会提问题？"); add_footer(s,11,2)

# Slide 12 experiment
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY); title(s,"安全小实验：糖块真的“消失”了吗？")
# two glasses
for i,label in enumerate(["不搅拌","轻轻搅拌"]):
    x=2.0+i*6.0
    glass=s.shapes.add_shape(MSO_SHAPE.TRAPEZOID,Inches(x),Inches(2.0),Inches(2.4),Inches(2.7)); glass.rotation=180; glass.fill.solid(); glass.fill.fore_color.rgb=RGBColor(210,240,255); glass.fill.transparency=25; glass.line.color.rgb=BLUE
    circle(s,x+0.85,2.5,0.5,WHITE,BLUE)
    add_text(s,"糖",x+0.94,2.62,0.3,0.25,15,NAVY,True,PP_ALIGN.CENTER)
    add_text(s,label,x-0.15,5.0,2.7,0.45,24,NAVY,True,PP_ALIGN.CENTER)
    if i==1:
        add_text(s,"↻",x+2.65,2.95,0.7,0.65,38,GREEN,True,PP_ALIGN.CENTER)
rounded(s,3.25,5.75,6.8,0.68,WHITE,BLUE)
add_text(s,"它没有消失，是变成很小的颗粒分散在水里。",3.55,5.87,6.2,0.42,20,DARK,True,PP_ALIGN.CENTER)
add_interaction(s,"预测：哪一杯糖块会更快看不见？为什么？"); add_footer(s,12,2)

# Slide 13 safety
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"最重要的一页：药不是糖果！",color=RED)
# traffic light
rounded(s,0.9,1.65,2.5,4.7,NAVY,NAVY)
circle(s,1.65,2.0,1.0,RED); circle(s,1.65,3.35,1.0,YELLOW); circle(s,1.65,4.7,1.0,GREEN)
safe_bullet(s,"不知道是什么药 → 不碰、不尝",4.1,1.9,RED)
safe_bullet(s,"身体不舒服 → 告诉老师和家长",4.1,2.85,GREEN)
safe_bullet(s,"药物只按大人/医生说的方法用",4.1,3.8,BLUE)
safe_bullet(s,"不和同学分享药物",4.1,4.75,PURPLE)
rounded(s,4.0,5.65,8.1,0.7,PINK,CORAL)
add_text(s,"记住：药物是帮助身体的工具，不是零食。",4.35,5.82,7.4,0.4,22,CORAL,True,PP_ALIGN.CENTER)
add_footer(s,13,2)

# Slide 14 traffic-light quiz
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM); title(s,"红灯还是绿灯？")
scenarios=[
("医生告诉妈妈：按时给你吃这个药。","绿灯",GREEN),
("看到桌上漂亮的药片，自己尝一个。","红灯",RED),
("同学说他的药很好吃，分给你一颗。","红灯",RED),
("肚子不舒服，马上告诉老师和爸爸妈妈。","绿灯",GREEN)]
for i,(txt,ans,c) in enumerate(scenarios):
    y=1.45+i*1.18
    rounded(s,0.75,y,9.5,0.88,WHITE,RGBColor(220,225,230))
    add_text(s,f"{i+1}. {txt}",1.0,y+0.12,8.9,0.58,19,DARK)
    rounded(s,10.55,y,1.85,0.88,RGBColor(245,245,245),c)
    add_text(s,ans,10.82,y+0.18,1.3,0.45,23,c,True,PP_ALIGN.CENTER)
add_interaction(s,"每题先不要看答案！全班用双手比“○”或“×”。"); add_footer(s,14,2)

# Slide 15 teammates
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY); title(s,"药物和身体，是队友！")
# team
add_character(s,1.0,2.6)
for i,c in enumerate([GREEN,PURPLE,BLUE]):
    circle(s,4.2+i*0.9,2.7,0.75,c)
add_text(s,"免疫小警察",4.0,3.75,2.6,0.45,21,PURPLE,True,PP_ALIGN.CENTER)
# helpers
helpers=[("睡觉",PURPLE),("喝水",BLUE),("吃好饭",GREEN),("休息",YELLOW)]
for i,(n,c) in enumerate(helpers):
    x=7.2+(i%2)*2.6; y=2.0+(i//2)*1.75
    rounded(s,x,y,2.2,1.25,WHITE,c); add_text(s,n,x+0.25,y+0.35,1.7,0.5,23,c,True,PP_ALIGN.CENTER)
add_text(s,"药物不是一个人战斗，身体自己也会努力恢复。",1.45,5.55,10.4,0.55,25,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,15,2)

# Slide 16 quiz
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE); title(s,"小侦探毕业考试！")
qs=[
("药物是不是糖果？","不是！"),
("不同的药，会不会有不同任务？","会！"),
("身体不舒服，能不能自己随便吃药？","不能！")]
for i,(q,a) in enumerate(qs):
    y=1.45+i*1.65
    circle(s,0.8,y,0.75,[CORAL,BLUE,GREEN][i])
    add_text(s,str(i+1),1.05,y+0.18,0.25,0.2,20,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,q,1.85,y-0.05,7.1,0.6,24,NAVY,True)
    rounded(s,9.35,y-0.05,2.4,0.75,CREAM,YELLOW)
    add_text(s,a,9.65,y+0.1,1.8,0.45,22,CORAL,True,PP_ALIGN.CENTER)
add_interaction(s,"先抢答，再一起大声说出正确答案！"); add_footer(s,16,2)

# Slide 17 recap
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM); title(s,"今天只要记住 3 句话")
messages=[("①","药物有不同任务",BLUE),("②","药物要找对目标",PURPLE),("③","药物一定要正确使用",GREEN)]
for i,(num,txt,c) in enumerate(messages):
    y=1.5+i*1.45
    rounded(s,2.2,y,8.9,1.05,WHITE,c)
    circle(s,2.55,y+0.18,0.68,c)
    add_text(s,num,2.68,y+0.34,0.4,0.25,19,WHITE,True,PP_ALIGN.CENTER)
    add_text(s,txt,3.6,y+0.25,6.8,0.5,27,NAVY,True,PP_ALIGN.CENTER)
add_text(s,"科学就是：保持好奇，勇敢提问，认真寻找答案。",1.4,5.95,10.6,0.5,23,CORAL,True,PP_ALIGN.CENTER)
add_footer(s,17,1)

# Slide 18
s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY)
add_text(s,"恭喜你！",0.7,0.8,11.9,0.65,40,NAVY,True,PP_ALIGN.CENTER)
add_text(s,"成为“安全用药小侦探”",0.7,1.55,11.9,0.65,32,CORAL,True,PP_ALIGN.CENTER)
# certificate
rounded(s,2.1,2.5,9.1,3.2,WHITE,YELLOW)
add_text(s,"小小科学家证书",3.2,2.9,6.9,0.55,31,PURPLE,True,PP_ALIGN.CENTER)
add_text(s,"我会好奇地提问，也会安全地对待药物。",3.0,3.75,7.3,0.55,22,DARK,False,PP_ALIGN.CENTER)
add_character(s,5.6,4.55)
add_text(s,"谢谢大家！",4.3,6.15,4.8,0.55,30,NAVY,True,PP_ALIGN.CENTER)
add_footer(s,18,1)



# ---------------- Real-world image mini-slides ----------------
ASSET_DIR = Path("chatgpt_exports/realistic_assets")
ASSET_DIR.mkdir(parents=True, exist_ok=True)

def download_asset(filename, url):
    path = ASSET_DIR / filename
    if path.exists() and path.stat().st_size > 10000:
        return path
    resp = requests.get(url, timeout=60, headers={"User-Agent":"Mozilla/5.0"})
    resp.raise_for_status()
    path.write_bytes(resp.content)
    return path

real_assets = {}
asset_specs = {
    "medicines": ("pills_medicines.jpg", "https://commons.wikimedia.org/wiki/Special:Redirect/file/Pills%20and%20medicines%2001.jpg"),
    "anatomy": ("human_anatomy.png", "https://commons.wikimedia.org/wiki/Special:Redirect/file/Human%20Anatomy.png"),
    "scientist": ("drug_synthesis.jpg", "https://commons.wikimedia.org/wiki/Special:Redirect/file/Drug%20synthesis.jpg"),
    "pipette": ("work_in_process.jpg", "https://commons.wikimedia.org/wiki/Special:Redirect/file/Work%20in%20process.jpg"),
    "caregiver": ("caregiver_medicine.jpg", "https://commons.wikimedia.org/wiki/Special:Redirect/file/Child%20receiving%20medicine%20from%20caregiver%20closeup.jpg"),
}
for key,(fn,url) in asset_specs.items():
    try:
        real_assets[key] = download_asset(fn,url)
    except Exception as e:
        print("asset download failed", key, e)

def add_photo_frame(slide, path, x,y,w,h):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    # thin white border by adding transparent rectangle around it
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x-0.03), Inches(y-0.03), Inches(w+0.06), Inches(h+0.06))
    frame.fill.background()
    frame.line.color.rgb = WHITE
    frame.line.width = Pt(2)
    # send frame behind picture if possible
    return pic

def add_credit(slide, txt):
    add_text(slide, txt, 0.55, 7.13, 12.2, 0.22, 8, GRAY, False, PP_ALIGN.RIGHT)

# Photo slide A: real medicines
if "medicines" in real_assets:
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM)
    title(s,"真实世界：药物到底长什么样？","看一看真实的药片、胶囊和包装")
    add_photo_frame(s,real_assets["medicines"],0.75,1.45,6.1,4.9)
    rounded(s,7.2,1.55,5.25,4.65,WHITE,BLUE)
    add_text(s,"你能找到这些吗？",7.55,1.85,4.6,0.5,26,NAVY,True,PP_ALIGN.CENTER)
    for j,(name,cx) in enumerate([("药片",BLUE),("胶囊",CORAL),("泡罩包装",GREEN),("药瓶",PURPLE)]):
        circle(s,7.65,2.75+j*0.75,0.42,cx)
        add_text(s,str(j+1),7.78,2.84+j*0.75,0.16,0.16,13,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,name,8.35,2.69+j*0.75,2.7,0.42,21,DARK,True)
    rounded(s,7.55,5.25,4.55,0.68,PINK,CORAL)
    add_text(s,"⚠ 再漂亮，也不是糖果！",7.8,5.38,4.05,0.35,20,CORAL,True,PP_ALIGN.CENTER)
    add_interaction(s,"照片里你最先看到了哪一种药物形式？")
    add_credit(s,"图片：Wikimedia Commons · Pills and medicines 01.jpg · CC BY-SA")

# Photo slide B: anatomy
if "anatomy" in real_assets:
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,WHITE)
    title(s,"真实世界：身体里的器官在哪里？","把“身体城市”变成一张真正的解剖图")
    add_photo_frame(s,real_assets["anatomy"],0.85,1.35,5.5,5.35)
    rounded(s,6.7,1.45,5.75,4.95,SKY,BLUE)
    organs=[("肺","帮助我们呼吸",BLUE),("心脏","把血液送到全身",CORAL),("胃","开始处理食物",YELLOW),("肠道","吸收营养",GREEN)]
    for j,(name,desc,col) in enumerate(organs):
        circle(s,7.05,1.95+j*0.95,0.52,col)
        add_text(s,str(j+1),7.21,2.07+j*0.95,0.2,0.18,14,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,name,7.8,1.85+j*0.95,1.0,0.4,21,col,True)
        add_text(s,desc,8.9,1.85+j*0.95,3.1,0.4,18,DARK)
    add_interaction(s,"我说器官名字，大家在图片上找一找！")
    add_credit(s,"图片：Wikimedia Commons · Human Anatomy.png · CC0")

# Photo slide C: scientists at work
if "scientist" in real_assets and "pipette" in real_assets:
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,SKY)
    title(s,"真实世界：药物科学家真的在做什么？","不是“魔法实验”，而是一遍遍认真测试")
    add_photo_frame(s,real_assets["scientist"],0.55,1.35,5.85,4.45)
    add_photo_frame(s,real_assets["pipette"],6.9,1.35,5.85,4.45)
    rounded(s,1.2,5.95,11.0,0.68,WHITE,BLUE)
    add_text(s,"观察 · 测量 · 移液 · 记录 · 比较 · 再改进",1.5,6.08,10.4,0.4,23,NAVY,True,PP_ALIGN.CENTER)
    add_interaction(s,"你能在照片里找到：白大褂、护目镜、移液器、实验仪器吗？")
    add_credit(s,"图片：Wikimedia Commons · Drug synthesis.jpg / Work in process.jpg · CC/PD")

# Photo slide D: caregiver safety
if "caregiver" in real_assets:
    s=prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s,CREAM)
    title(s,"真实世界：吃药为什么需要大人帮助？")
    add_photo_frame(s,real_assets["caregiver"],0.75,1.45,6.25,4.9)
    rounded(s,7.35,1.55,5.0,4.65,WHITE,GREEN)
    add_text(s,"安全用药 3 步",7.7,1.85,4.3,0.48,27,GREEN,True,PP_ALIGN.CENTER)
    rules=[("①","身体不舒服先告诉大人",GREEN),("②","只用大人/医生确认的药",BLUE),("③","按正确的量和方法使用",PURPLE)]
    for j,(n,txt,col) in enumerate(rules):
        circle(s,7.75,2.75+j*0.95,0.58,col)
        add_text(s,n,7.91,2.9+j*0.95,0.25,0.2,16,WHITE,True,PP_ALIGN.CENTER)
        add_text(s,txt,8.55,2.7+j*0.95,3.2,0.5,20,DARK,True)
    rounded(s,7.75,5.55,4.15,0.5,PINK,CORAL)
    add_text(s,"药不是自己决定吃的东西",7.95,5.62,3.75,0.33,18,CORAL,True,PP_ALIGN.CENTER)
    add_interaction(s,"如果你在家里看到不认识的药，第一件事应该做什么？")
    add_credit(s,"图片：Wikimedia Commons · Child receiving medicine from caregiver closeup.jpg · CC BY 2.0")

# Move the four real-world slides into the story flow.
# They are appended at the end; reposition them after base slides 2, 3, 11 and 13.
def move_last_to(index):
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(index, el)

# Because moving the last slide changes which remains last, move in reverse desired order.
# Current appended order A,B,C,D; D is last.
move_last_to(13)  # D after original safety slide 13
move_last_to(11)  # C after scientist-work slide 11
move_last_to(3)   # B after body-city slide 3
move_last_to(2)   # A after dosage-forms slide 2

# Metadata
prs.core_properties.title="药物小侦探：一颗药是怎样帮助身体的？"
prs.core_properties.subject="一年级家长进课堂科普课，30分钟"
prs.core_properties.author="ChatGPT"

prs.save(OUT)
# QA reopen
check=Presentation(OUT)
assert len(check.slides)>=22
assert OUT.stat().st_size>50000
print(OUT)
print("slides",len(check.slides),"bytes",OUT.stat().st_size)
