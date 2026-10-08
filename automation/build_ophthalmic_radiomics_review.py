from pathlib import Path
import math, os
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Circle, Ellipse, Rectangle, FancyArrowPatch, Wedge
from matplotlib.lines import Line2D
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor as DocxRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTDIR=Path("chatgpt_exports/ophthalmic_radiomics_assets")
OUTDIR.mkdir(parents=True, exist_ok=True)
DOCX=Path("chatgpt_exports/影像组学与眼科疾病_高影响力综述框架_图表整合版.docx")

# ---------- figure helpers ----------
COLORS={
"blue":"#4D91E2","navy":"#244A72","sky":"#DDEEFF","green":"#66B77A","mint":"#E2F4E8",
"coral":"#E9877B","pink":"#F8E4E6","orange":"#F2B66D","cream":"#FFF5E7","purple":"#8E79C6",
"lav":"#EEE9FA","gray":"#647484","dark":"#25323D","yellow":"#F2D06B","white":"#FFFFFF","red":"#D95D5D"
}
plt.rcParams.update({"font.size":9,"font.family":"DejaVu Sans"})

def setup(title):
    fig,ax=plt.subplots(figsize=(14,8),dpi=170)
    fig.patch.set_facecolor("white"); ax.set_facecolor("white")
    ax.set_xlim(0,14); ax.set_ylim(0,8); ax.axis("off")
    ax.text(0.3,7.6,title,fontsize=20,fontweight="bold",color=COLORS["navy"],va="top")
    return fig,ax

def box(ax,x,y,w,h,text,fc="#F7FAFD",ec="#B8C7D5",fs=10,bold=False,align="center",radius=0.15):
    p=FancyBboxPatch((x,y),w,h,boxstyle=f"round,pad=0.03,rounding_size={radius}",facecolor=fc,edgecolor=ec,linewidth=1.2)
    ax.add_patch(p)
    ax.text(x+w/2,y+h/2,text,ha=align,va="center",fontsize=fs,fontweight="bold" if bold else "normal",color=COLORS["dark"],wrap=True)
    return p

def arr(ax,x1,y1,x2,y2,c=None,lw=2):
    a=FancyArrowPatch((x1,y1),(x2,y2),arrowstyle="-|>",mutation_scale=14,linewidth=lw,color=c or COLORS["blue"])
    ax.add_patch(a); return a

def eye(ax,cx,cy,scale=1):
    ax.add_patch(Ellipse((cx,cy),3.2*scale,2.0*scale,facecolor="#F8FBFF",edgecolor=COLORS["navy"],lw=1.5))
    ax.add_patch(Circle((cx,cy),0.7*scale,facecolor="#A9D8F8",edgecolor=COLORS["blue"],lw=1))
    ax.add_patch(Circle((cx,cy),0.28*scale,facecolor=COLORS["navy"],edgecolor="none"))
    ax.add_patch(Wedge((cx+0.85*scale,cy),1.2*scale,80,280,width=0.17*scale,facecolor="#F4B0A7",edgecolor="none"))
    arr(ax,cx+1.55*scale,cy,cx+2.6*scale,cy,COLORS["orange"],1.8)

def save(fig,name):
    path=OUTDIR/name
    fig.savefig(path,bbox_inches="tight",facecolor="white")
    plt.close(fig); return path

# Fig1 workflow
fig,ax=setup("Figure 1 | Ophthalmic radiomics: from imaging to clinical decision")
stages=[
("1  Imaging","Fundus · OCT · OCTA · FAF\nUltrasound · CT · MRI",COLORS["sky"]),
("2  Preprocess","QC · registration · denoise\nresample · harmonize",COLORS["cream"]),
("3  ROI","Retinal layers · lesions\noptic nerve · orbit · tumor",COLORS["mint"]),
("4  Features","Intensity · shape · texture\nwavelet · spatial · delta",COLORS["lav"]),
("5  Modeling","feature selection\nML / deep radiomics",COLORS["pink"]),
("6  Testing","patient-level split\nexternal + prospective test",COLORS["sky"]),
("7  Utility","diagnosis · prognosis\nresponse · decision support",COLORS["mint"])]
x=0.35
for i,(t,b,c) in enumerate(stages):
    box(ax,x,4.65,1.72,1.8,t+"\n\n"+b,c,"#C2CFD9",9,i==0)
    if i<6: arr(ax,x+1.72,5.55,x+1.97,5.55,COLORS["gray"],1.5)
    x+=1.95
eye(ax,2.0,2.25,0.75)
box(ax,4.4,1.45,2.2,1.55,"Feature map\n\nHeterogeneity\nvascularity\nmorphology",COLORS["cream"],COLORS["orange"],10,True)
arr(ax,3.35,2.25,4.35,2.25,COLORS["orange"])
box(ax,7.05,1.45,2.2,1.55,"Imaging phenotype\n\nquantitative,\nreproducible,\nlongitudinal",COLORS["lav"],COLORS["purple"],10,True)
arr(ax,6.65,2.25,7.0,2.25,COLORS["purple"])
box(ax,9.7,1.45,3.5,1.55,"Clinical question\n\nWho has disease?\nWho will progress?\nWho will respond?",COLORS["mint"],COLORS["green"],10,True)
arr(ax,9.3,2.25,9.65,2.25,COLORS["green"])
fig1=save(fig,"figure1_workflow.png")

# Fig2 phenotype map
fig,ax=setup("Figure 2 | Multimodal computable ocular phenotype map")
eye(ax,7.0,4.15,1.35)
regions=[
(0.6,5.0,3.6,1.7,"Retina / macula","OCT: layer architecture, fluid, drusen\nOCTA: capillary density, FAZ, tortuosity\nFundus: lesion texture, pigmentation",COLORS["pink"],COLORS["coral"]),
(0.6,2.4,3.6,1.7,"Choroid","Thickness · vascularity · texture\ninflammation / fibrosis proxies",COLORS["lav"],COLORS["purple"]),
(9.8,5.0,3.6,1.7,"Optic nerve / disc","RNFL / GCC · excavation\nperipapillary texture · perfusion\nlongitudinal change",COLORS["mint"],COLORS["green"]),
(9.8,2.4,3.6,1.7,"Orbit / adnexa","EOM · lacrimal gland · orbital fat\noptic nerve · perilesional tissue\nMRI/CT multiparametric texture",COLORS["sky"],COLORS["blue"]),
(5.15,0.55,3.7,1.25,"Intraocular tumor","shape · internal heterogeneity · enhancement · genotype-related phenotype",COLORS["cream"],COLORS["orange"])
]
for x,y,w,h,t,b,fc,ec in regions:
    box(ax,x,y,w,h,t+"\n\n"+b,fc,ec,9,True)
for p1,p2 in [((4.2,5.8),(5.7,4.7)),((4.2,3.2),(5.7,3.8)),((9.8,5.8),(8.3,4.75)),((9.8,3.2),(8.4,3.7)),((7.0,1.8),(7.0,2.8))]:
    arr(ax,*p1,*p2,COLORS["gray"],1.3)
fig2=save(fig,"figure2_phenotype_map.png")

# Fig3 retinal disease
fig,ax=setup("Figure 3 | Retinal radiomics: phenotype-grounded analysis in AMD, DR and DME")
panels=[
(0.5,"AMD","drusen substructure\nRPE–Bruch's complex\nGA conversion / growth","OCT · FAF",COLORS["cream"],COLORS["orange"]),
(4.8,"Diabetic retinopathy","microvascular topology\ncapillary texture\nischemia / leakage phenotype","OCTA · fundus · FFA",COLORS["pink"],COLORS["coral"]),
(9.1,"Diabetic macular edema","fluid geometry\nretinal-layer texture\ninflammatory phenotype\nresponse prediction","OCT ± OCTA",COLORS["sky"],COLORS["blue"])]
for x,t,b,m,fc,ec in panels:
    box(ax,x,4.25,3.8,2.2,t+"\n\n"+b+"\n\n"+m,fc,ec,10,True)
    # synthetic image patch
    rng=np.random.default_rng(int(x*100+7))
    mat=rng.normal(size=(18,32))
    ax.imshow(mat,cmap="gray",extent=(x+0.35,x+3.45,2.2,3.55),aspect="auto")
    for k in range(4):
        ax.plot([x+0.4,x+3.35],[2.35+k*0.25,2.42+k*0.25],color=[COLORS["coral"],COLORS["orange"],COLORS["green"],COLORS["purple"]][k],lw=1)
box(ax,1.3,0.65,11.4,0.95,"Core shift: from visible lesions → quantitative phenotype → longitudinal risk → treatment response",COLORS["mint"],COLORS["green"],12,True)
fig3=save(fig,"figure3_retinal_radiomics.png")

# Fig4 glaucoma
fig,ax=setup("Figure 4 | Glaucoma radiomics: structure–perfusion–time integration")
# optic disc
ax.add_patch(Circle((2.0,4.25),1.2,facecolor="#F2B6A6",edgecolor=COLORS["coral"],lw=2))
ax.add_patch(Circle((2.0,4.25),0.55,facecolor="#FCE6D9",edgecolor=COLORS["orange"],lw=1.5))
for ang in np.linspace(0,2*np.pi,14,endpoint=False):
    ax.plot([2,2+1.55*np.cos(ang)],[4.25,4.25+1.2*np.sin(ang)],color=COLORS["red"],lw=1)
box(ax,0.55,1.0,2.9,1.35,"Optic disc phenotype\ncup/disc · rim texture\nlamina cribrosa",COLORS["cream"],COLORS["orange"],9,True)
# three streams
box(ax,4.1,5.25,2.5,1.2,"Structural OCT\nRNFL · GCC\nlayer texture",COLORS["sky"],COLORS["blue"],10,True)
box(ax,4.1,3.45,2.5,1.2,"OCTA perfusion\nvessel density\nFAZ · tortuosity",COLORS["mint"],COLORS["green"],10,True)
box(ax,4.1,1.65,2.5,1.2,"Functional / time\nvisual field\nlongitudinal OCT",COLORS["lav"],COLORS["purple"],10,True)
for yy in [5.85,4.05,2.25]: arr(ax,3.3,4.25,4.0,yy,COLORS["gray"],1.3)
box(ax,7.45,3.25,2.7,1.9,"Multimodal phenotype\n\nstructure + perfusion\n+ trajectory",COLORS["pink"],COLORS["coral"],11,True)
for yy in [5.85,4.05,2.25]: arr(ax,6.65,yy,7.4,4.2,COLORS["coral"],1.3)
box(ax,11.05,4.55,2.2,1.2,"Early detection",COLORS["mint"],COLORS["green"],10,True)
box(ax,11.05,3.0,2.2,1.2,"Progression risk",COLORS["cream"],COLORS["orange"],10,True)
box(ax,11.05,1.45,2.2,1.2,"Treatment intensity",COLORS["sky"],COLORS["blue"],10,True)
for yy in [5.15,3.6,2.05]: arr(ax,10.2,4.2,11.0,yy,COLORS["navy"],1.4)
fig4=save(fig,"figure4_glaucoma.png")

# Fig5 oncology/orbit
fig,ax=setup("Figure 5 | Ocular oncology and orbital radiomics")
cards=[
(0.55,4.55,3.1,1.65,"Uveal melanoma","MRI / CT / US\nshape · texture · enhancement\nmetastatic risk · eye preservation",COLORS["cream"],COLORS["orange"]),
(3.95,4.55,3.1,1.65,"Thyroid eye disease","MRI\nEOM + fat + lacrimal + ON\nactivity · DON · steroid response",COLORS["sky"],COLORS["blue"]),
(7.35,4.55,3.1,1.65,"Orbital lymphoma / IOI","MRI / CT\nintralesional + perilesional texture\ndifferential diagnosis",COLORS["pink"],COLORS["coral"]),
(10.75,4.55,2.7,1.65,"IgG4-ROD","MRI\nmulti-region heterogeneity\ndifferential diagnosis",COLORS["lav"],COLORS["purple"])]
for x,y,w,h,t,b,fc,ec in cards: box(ax,x,y,w,h,t+"\n\n"+b,fc,ec,9,True)
# central continuum
box(ax,1.0,1.15,2.6,1.45,"Diagnosis\nlesion identity",COLORS["mint"],COLORS["green"],10,True)
box(ax,4.05,1.15,2.6,1.45,"Risk\nbiological aggressiveness",COLORS["cream"],COLORS["orange"],10,True)
box(ax,7.1,1.15,2.6,1.45,"Response\ntherapy sensitivity",COLORS["sky"],COLORS["blue"],10,True)
box(ax,10.15,1.15,2.6,1.45,"Planning\nsurgery / radiation",COLORS["lav"],COLORS["purple"],10,True)
for x in [3.6,6.65,9.7]: arr(ax,x,1.88,x+0.4,1.88,COLORS["gray"],1.5)
ax.text(7,3.35,"Radiomics converts tissue heterogeneity into quantitative clinical phenotypes",ha="center",va="center",fontsize=13,fontweight="bold",color=COLORS["navy"])
fig5=save(fig,"figure5_oncology_orbit.png")

# Fig6 multimodal fusion
fig,ax=setup("Figure 6 | Multimodal fusion and radiogenomics for precision ophthalmology")
inputs=[
(0.35,5.15,2.0,1.15,"Radiomics","OCT · OCTA · MRI\ntexture · shape",COLORS["sky"],COLORS["blue"]),
(0.35,3.65,2.0,1.15,"Clinical","age · VA · IOP\nsymptoms · therapy",COLORS["cream"],COLORS["orange"]),
(0.35,2.15,2.0,1.15,"Molecular","genomics\ntranscriptomics\nproteomics",COLORS["lav"],COLORS["purple"]),
(0.35,0.65,2.0,1.15,"Longitudinal","baseline → follow-up\ndelta-radiomics",COLORS["mint"],COLORS["green"])]
for x,y,w,h,t,b,fc,ec in inputs:
    box(ax,x,y,w,h,t+"\n"+b,fc,ec,9,True)
    arr(ax,2.4,y+0.58,4.2,3.95,ec,1.3)
box(ax,4.25,2.45,3.0,2.85,"Fusion layer\n\nfeature-level\nrepresentation-level\ndecision-level\n\n+ uncertainty",COLORS["pink"],COLORS["coral"],11,True)
arr(ax,7.3,3.88,8.2,3.88,COLORS["coral"],1.8)
box(ax,8.25,4.8,2.3,1.35,"Interpretable\nendotypes","subtype A · B · C\nbiologically grounded",COLORS["cream"],COLORS["orange"],10,True)
box(ax,8.25,2.75,2.3,1.35,"Radiogenomics","imaging ↔ genotype\nimaging ↔ pathway",COLORS["lav"],COLORS["purple"],10,True)
box(ax,8.25,0.7,2.3,1.35,"Digital biomarker","risk · response\nmonitoring",COLORS["sky"],COLORS["blue"],10,True)
for yy in [5.47,3.42,1.37]: arr(ax,10.6,yy,11.35,3.85,COLORS["gray"],1.2)
box(ax,11.4,2.4,2.2,2.9,"Precision care\n\nDiagnosis\nPrognosis\nTreatment selection\nMonitoring\nClinical trials",COLORS["mint"],COLORS["green"],10,True)
fig6=save(fig,"figure6_multimodal_radiogenomics.png")

# Fig7 translation
fig,ax=setup("Figure 7 | Roadmap from proof-of-concept radiomics to clinical deployment")
steps=[
("1  Discovery","single-center\nexploratory features",COLORS["cream"],COLORS["orange"]),
("2  Standardize","IBSI-aligned features\nvendor / protocol QC",COLORS["sky"],COLORS["blue"]),
("3  Test","patient-level split\nexternal testing",COLORS["mint"],COLORS["green"]),
("4  Generalize","multi-center\ncross-device / population",COLORS["lav"],COLORS["purple"]),
("5  Utility","calibration · DCA\nworkflow impact",COLORS["pink"],COLORS["coral"]),
("6  Prospective","silent deployment\nprospective evaluation",COLORS["sky"],COLORS["blue"]),
("7  Implementation","regulation · fairness\nmonitoring / updating",COLORS["mint"],COLORS["green"])]
x=0.35
for i,(t,b,fc,ec) in enumerate(steps):
    box(ax,x,4.45,1.72,1.8,t+"\n\n"+b,fc,ec,8.8,True)
    if i<6: arr(ax,x+1.72,5.35,x+1.94,5.35,COLORS["gray"],1.3)
    x+=1.93
# failure modes
fails=[
("Eye-level leakage","Both eyes from one patient\nmust not cross data splits"),
("Device shift","OCT vendor / scan protocol\nchanges feature distribution"),
("Segmentation drift","automated ROI errors\npropagate to features"),
("Small cohorts","high-dimensional features\n→ overfitting"),
("Opaque utility","high AUC ≠ clinical benefit")]
for i,(t,b) in enumerate(fails):
    xx=0.7+i*2.6
    box(ax,xx,1.25,2.2,1.35,t+"\n"+b,"#FFF1F1",COLORS["red"],8.4,True)
ax.text(7,3.35,"Translation requires reproducibility, fairness, calibration and clinical utility — not accuracy alone.",ha="center",fontsize=12.5,fontweight="bold",color=COLORS["navy"])
fig7=save(fig,"figure7_translation.png")

# ---------- DOCX helpers ----------
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_font(run,size=10.5,bold=False,color=None,east="宋体",latin="Times New Roman"):
    run.font.name=latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'),east)
    run.font.size=Pt(size); run.bold=bold
    if color: run.font.color.rgb=DocxRGB.from_string(color)

def add_p(doc,text="",size=10.5,bold=False,center=False,indent=True,space=4):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(space)
    if indent: p.paragraph_format.first_line_indent=Cm(0.74)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); set_font(r,size,bold)
    return p

def h(doc,text,level=1):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(10 if level==1 else 6)
    p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text)
    if level==1: set_font(r,15,True,"17365D","黑体")
    elif level==2: set_font(r,12.5,True,"245B78","黑体")
    else: set_font(r,11,True,"3B6A57","黑体")
    return p

def bullets(doc,items):
    for item in items:
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Cm(0.6); p.paragraph_format.first_line_indent=Cm(-0.3); p.paragraph_format.space_after=Pt(2)
        r=p.add_run("• "+item); set_font(r,10.2)

def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,x in enumerate(headers):
        c=t.rows[0].cells[j]; set_cell_shading(c,"DCEAF5"); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run(x); set_font(r,9,True,"17365D","黑体")
    for row in rows:
        cells=t.add_row().cells
        for j,x in enumerate(row):
            cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(0)
            r=p.add_run(str(x)); set_font(r,8.5)
    doc.add_paragraph(); return t

doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.0); sec.right_margin=Cm(2.0)
style=doc.styles["Normal"]; style.font.name="Times New Roman"; style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); style.font.size=Pt(10.5)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("影像组学与眼科疾病"); set_font(r,20,True,"17365D","黑体")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("从可计算影像表型到精准眼科与临床转化"); set_font(r,14,True,"245B78","黑体")
add_p(doc,"Ophthalmic radiomics in precision eye care: from computable imaging phenotypes to multimodal biomarkers and clinical translation",11,False,True,False,5)
add_p(doc,"定位：IF >10 / CNS子刊风格综述框架；含7幅机制/概念图与4个原生可编辑表格。证据更新至2026年10月。",9.5,False,True,False,10)

doc.add_page_break()
h(doc,"1. 综述定位：为什么不能再写成“radiomics + 疾病罗列”",1)
add_p(doc,"2025年发表于 European Radiology 的系统综述已纳入41篇眼科影像组学研究，覆盖5类疾病与7种影像模态；模型AUC多数落在0.7–1.0，但平均Radiomics Quality Score仅11.17/36，主要问题是回顾性单中心设计和外部测试不足。因此，再写一篇按AMD、DR、青光眼、眼肿瘤逐项罗列的综述，很难形成高影响力差异化。")
add_p(doc,"本综述建议把“影像组学”定义为眼科影像从可视征象转化为可计算表型的一层基础设施，并围绕‘表型—生物学—时间—治疗—转化’五个层级组织全文。真正的核心问题不是模型能否分类，而是：这些特征能否稳定表征眼组织的结构、微血管、炎症、纤维化和肿瘤异质性，并在跨设备、跨中心、跨人群中支持可解释的临床决策。")
bullets(doc,[
"核心转变1：从“feature hunting”转向 biologically grounded imaging phenotypes。",
"核心转变2：从单模态分类转向 OCT/OCTA/FAF/FP/US/MRI/CT 与临床、组学、纵向数据融合。",
"核心转变3：从横断面诊断转向 progression、treatment response、delta-radiomics 与临床试验终点。",
"核心转变4：从高AUC转向 external testing、calibration、decision-curve、fairness 与真实工作流价值。",
"核心转变5：突出眼科特异方法学问题：双眼相关性、患者级数据切分、OCT设备/扫描协议漂移和自动分割误差。"
])

h(doc,"2. 推荐题目",1)
bullets(doc,[
"首选：From pixels to phenotypes: radiomics and multimodal imaging intelligence in precision ophthalmology",
"备选：Ophthalmic radiomics in precision eye care: quantitative phenotypes, multimodal fusion and clinical translation",
"备选：Radiomics across the eye: from retinal microstructure to orbital and ocular tumor phenotypes",
"中文：从像素到表型：影像组学驱动的精准眼科疾病分层与临床转化"
])

h(doc,"3. 全文唯一主线",1)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("multimodal ocular imaging → quantitative phenotype → disease endotype → longitudinal risk / response → precision decision → clinical utility")
set_font(r,12,True,"8A3150")

h(doc,"4. 推荐正文框架",1)
outline=[
("4.1 Introduction: ophthalmology as a natural laboratory for quantitative imaging",[
"眼科具有高分辨率、重复性强、纵向随访密集和多模态影像并存的天然优势。",
"区分radiomics、deep radiomics、multimodal AI和oculomics：前者聚焦高通量影像特征，oculomics更强调由眼部数据推断全身健康。"]),
("4.2 From images to computable phenotypes: the radiomics workflow",[
"采集、质量控制、配准、标准化、ROI/VOI分割、特征提取、特征选择、建模、内部测试、外部测试和临床部署。",
"特征层级：first-order、shape、texture、wavelet/filter、spatial features、delta-radiomics、deep features。",
"强调IBSI一致性与特征可重复性，而不是盲目扩大特征数量。"]),
("4.3 A multimodal map of ophthalmic radiomics",[
"Retina/macula：fundus、OCT、OCTA、FAF、FFA。",
"Optic nerve/glaucoma：disc photography、OCT RNFL/GCC、OCTA perfusion、visual field。",
"Whole eye/ocular oncology：US、MRI、CT。",
"Orbit/adnexa：multiparametric MRI/CT，结合EOM、lacrimal gland、fat、optic nerve与perilesional tissue。"]),
("4.4 Retina: radiomics as a quantitative disease-phenotyping layer",[
"AMD：drusen substructure、RPE–Bruch复合体、GA转换与生长；避免只做二分类。",
"DR/DME：OCTA微血管纹理、血管密度/迂曲度、缺血与渗漏代理表型，OCT液体和视网膜层纹理。",
"强调影像组学应与可解释的OCT/OCTA生物标志物结合。"]),
("4.5 Glaucoma and optic neuropathy: structure–perfusion–time integration",[
"整合optic disc、RNFL/GCC、lamina cribrosa、OCTA灌注和视觉功能。",
"重点转向早期检测、真正进展与年龄变化区分、纵向风险和治疗强度分层。"]),
("4.6 Ocular oncology and orbital disease: where classical radiomics is most mature",[
"Uveal melanoma：鉴别诊断、风险分层、预后、眼球保留、放疗计划与治疗反应。",
"Thyroid eye disease：whole-orbit multi-region radiomics，预测活动度、DON和糖皮质激素反应。",
"Orbital lymphoma / idiopathic orbital inflammation / IgG4-related ophthalmic disease：intralesional + perilesional radiomics用于鉴别诊断。"]),
("4.7 Beyond single-modality radiomics: multimodal fusion and radiogenomics",[
"影像 + 临床 + 纵向数据 + genomics/transcriptomics/proteomics/metabolomics。",
"从预测分数升级为影像内表型，并探索影像—分子通路和影像—基因型映射。",
"区分早期融合、表征层融合和决策层融合。"]),
("4.8 Longitudinal radiomics and treatment-response phenotyping",[
"delta-radiomics、动态OCT/OCTA特征、治疗前后MRI异质性变化。",
"强调response prediction、adaptive treatment和影像替代终点的可能性。"]),
("4.9 Translation: from high AUC to trustworthy clinical utility",[
"外部测试、多中心/跨设备测试、校准、decision-curve analysis和silent deployment。",
"公平性与代表性：设备、种族/人群、年龄、病程和中心差异。",
"报告与质量：IBSI、CLAIM 2024、TRIPOD+AI、PROBAST+AI。"]),
("4.10 Future directions",[
"foundation models与self-supervised learning结合handcrafted radiomics。",
"federated learning、privacy-preserving multi-center radiomics、synthetic data。",
"3D whole-eye digital phenotype / digital twin。",
"从相关性特征向可干预、可解释、可复现的causal imaging biomarker转变。"])
]
for title,bs in outline:
    h(doc,title,2); bullets(doc,bs)

h(doc,"5. 7幅主图设计",1)
figs=[
(fig1,"Figure 1 | Ophthalmic radiomics: from imaging to clinical decision","以眼科影像组学全流程为总览，强调最终目标不是AUC，而是形成稳定、可重复、可解释的量化影像表型并进入临床决策。"),
(fig2,"Figure 2 | Multimodal computable ocular phenotype map","以解剖结构为中心，而非疾病为中心，展示retina/macula、choroid、optic nerve、orbit/adnexa和intraocular tumor的可计算表型。"),
(fig3,"Figure 3 | Retinal radiomics in AMD, DR and DME","展示从OCT/OCTA可见病灶向纹理、微血管和结构异质性量化的转变，并连接风险与治疗反应。"),
(fig4,"Figure 4 | Glaucoma radiomics: structure–perfusion–time integration","整合结构OCT、OCTA灌注、视觉功能和纵向变化，突出进展预测而非单次分类。"),
(fig5,"Figure 5 | Ocular oncology and orbital radiomics","比较uveal melanoma、TED、orbital lymphoma/IOI和IgG4-ROD，展示传统MRI/CT radiomics最有优势的组织异质性与多区域分析。"),
(fig6,"Figure 6 | Multimodal fusion and radiogenomics","将影像组学、临床、分子组学和纵向数据融合成机制内表型与精准治疗决策。"),
(fig7,"Figure 7 | Clinical translation roadmap","系统展示从proof-of-concept到多中心标准化、外部测试、临床效用、前瞻性实施的路线，并特别标注眼科常见数据泄漏和设备漂移问题。")
]
for path,title_txt,legend in figs:
    doc.add_picture(str(path),width=Inches(6.55)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_p(doc,title_txt,9.5,True,True,False,2)
    add_p(doc,"图注："+legend,9.2,False,False,False,8)

doc.add_page_break()
h(doc,"6. 4个原生可编辑表格",1)

h(doc,"Table 1 | 眼科疾病–影像模态–影像组学任务矩阵",2)
table(doc,["疾病/领域","主要影像模态","最有价值的ROI/表型","适合的临床任务","写作重点"],[
["AMD / GA","OCT, FAF, CFP","drusen、RPE–Bruch复合体、GA边缘、choroid","进展风险、GA转换/生长、治疗反应","表型驱动，而非简单AMD二分类"],
["DR / DME","OCTA, OCT, CFP, FFA","SCP/DCP、FAZ、毛细血管纹理、液体、视网膜层","筛查、分期、缺血/渗漏、治疗反应","微血管+结构联合"],
["Glaucoma","OCT, OCTA, fundus, VF","RNFL/GCC、disc、lamina、peripapillary perfusion","早期诊断、进展预测、治疗强度","structure–function–time"],
["Uveal melanoma","MRI, CT, US, fundus","tumor/peritumor、shape、texture、enhancement","鉴别、预后、眼球保留、治疗反应","radiogenomics / planning潜力"],
["Thyroid eye disease","MRI, CT","EOM、fat、lacrimal gland、optic nerve","活动度、DON、IVGC response","whole-orbit multi-region radiomics"],
["Orbital lymphoma / IOI / IgG4-ROD","MRI, CT","lesion + perilesional tissue","鉴别诊断、风险与治疗规划","perilesional radiomics与外部测试"]
])

h(doc,"Table 2 | Radiomics特征类别、可能的生物学含义与眼科示例",2)
table(doc,["特征类别","量化对象","潜在生物学含义","眼科示例","主要风险"],[
["First-order","强度/灰度分布","组织信号组成与均一性","OCT reflectivity；MRI signal heterogeneity","受设备和预处理影响大"],
["Shape","面积、体积、球形度、边界","病灶生长模式与结构重塑","drusen/tumor morphology","依赖精确分割"],
["Texture","GLCM/GLRLM/GLSZM等","微结构异质性","tumor heterogeneity；retinal texture","可解释性有限、易过拟合"],
["Wavelet/filter","多尺度频率特征","不同空间尺度的纹理","MRI/OCT filtered features","跨平台重现性"],
["Vascular radiomics","密度、迂曲、分形/纹理","微循环损伤与灌注表型","OCTA in DR/glaucoma","分割与阈值敏感"],
["Spatial/perilesional","区域关系与邻域特征","病灶—微环境互作","orbital lymphoma vs inflammation","ROI定义尚未标准化"],
["Delta-radiomics","时间变化率","治疗反应与疾病轨迹","UM/TED/OCT follow-up","需要稳定纵向协议"],
["Deep radiomics","网络学习表征","高阶潜在表型","multimodal ophthalmic AI","可解释性与域漂移"]
])

h(doc,"Table 3 | 代表性眼科影像组学证据锚点",2)
table(doc,["研究","疾病/模态","样本/设计","主要结果","对综述的意义"],[
["Zhang et al., Eur Radiol 2025","眼科radiomics系统综述","41篇研究；5类疾病；7种模态","AUC多数0.7–1.0；平均RQS 11.17/36","证明领域已有规模，但临床转化质量不足"],
["Carrera-Escalé et al., Ophthalmol Sci","DR；OCT/OCTA/FP","726眼，439人","DM: OCT AUC 0.82；DR: OCTA 0.77；referable DR: OCTA 0.87","视网膜微血管radiomics的代表性概念验证"],
["Perkins et al., Diagnostics 2025","AMD；OCT drusen","多类drusen与743 OCT scans","drusen分类AUC 0.87–0.95；GA风险预测0.59–0.74","强调phenotype-grounded radiomics而非单纯分类"],
["Su et al., Eur J Radiol 2020","uveal melanoma；MRI","245例","T2WI+CET1WI模型AUC约0.87–0.88","眼肿瘤MRI radiomics的早期高质量例证"],
["Zhang et al., J Transl Med 2024","TED；whole-orbit MRI","127例IVGC治疗患者","多器官/融合区域radiomics用于疗效预测","从single-ROI升级到whole-orbit phenotype"],
["Yedekci et al., Strahlenther Onkol 2026","uveal melanoma；CT+MRI","308例，≥5年随访","预测secondary enucleation AUC 0.90","连接影像异质性与长期临床决策"]
])

h(doc,"Table 4 | 投稿级质量与转化清单：眼科radiomics最容易被审稿人质疑的环节",2)
table(doc,["环节","必须回答的问题","眼科特异风险","建议标准/做法"],[
["数据切分","训练、调参、内部测试、外部测试如何划分？","同一患者双眼进入不同集合造成泄漏","必须patient-level split；报告每层级样本量"],
["采集协议","设备和扫描参数是否一致？","OCT vendor、scan size、signal strength差异明显","多设备测试；协议/设备分层分析；harmonization"],
["ROI/分割","ROI由谁、如何定义？","自动分层/液体/肿瘤分割误差会传递","报告ICC/Dice；敏感性分析；人工复核策略"],
["特征计算","特征定义是否可重复？","不同软件/重采样/离散化导致漂移","优先IBSI-compliant实现并公开参数"],
["建模","维度是否远大于事件数？","小样本、高维、重复眼数据","嵌套CV；正则化；避免数据泄漏"],
["性能","是否只报告AUC？","高AUC不等于临床可用","calibration、CI、threshold metrics、DCA"],
["泛化","是否有真正外部测试？","单中心设备特征被模型学习","跨中心/跨设备/跨人群external testing"],
["公平性","不同年龄/种族/设备是否公平？","数据库代表性不足","预设亚组与fairness分析"],
["报告","研究过程是否透明？","“validation”术语含混","CLAIM 2024；TRIPOD+AI；PROBAST+AI；RQS"],
["实施","模型是否改变决策和结局？","缺少workflow与prospective evidence","silent deployment→前瞻评估→影响研究"]
])

h(doc,"7. 这篇综述最值得强调的6个创新概念",1)
bullets(doc,[
"Computable ocular phenotype：把传统影像征象转化为连续量化表型。",
"Anatomy-first radiomics：按retina–choroid–optic nerve–orbit–tumor组织，而不是单纯按疾病目录。",
"Phenotype-grounded radiomics：特征必须与临床可识别结构/病理过程建立联系。",
"Temporal radiomics：眼科高频随访是其他器官难以复制的优势，应充分利用delta-radiomics。",
"Radiogenomics / multimodal fusion：连接影像异质性与分子亚型，而不是把AI停留在分类层。",
"Trustworthy translation：将IBSI、CLAIM 2024、TRIPOD+AI、PROBAST+AI与眼科双眼数据泄漏问题纳入同一转化框架。"
])

h(doc,"8. 结论段建议",1)
add_p(doc,"眼科影像组学已经从早期的纹理特征与传统机器学习，逐渐扩展为涵盖OCT/OCTA微结构、MRI/CT组织异质性、多区域空间特征、纵向变化和多模态融合的计算影像学体系。其真正价值并不在于不断提高单中心测试集AUC，而在于构建可重复、可解释且能够改变临床决策的定量影像表型。未来高影响力工作应进一步连接影像表型与疾病生物学、分子机制和治疗反应，并通过跨设备、多中心、前瞻性评估证明其临床效用。")
add_p(doc,"因此，本综述建议把眼科radiomics重新定义为precision ophthalmology的‘phenotype layer’：它位于原始影像和临床决策之间，并可向下连接组织和分子生物学，向上连接风险预测、治疗选择和真实世界实施。")

h(doc,"9. 关键参考文献与方法学标准",1)
refs=[
"Zhang H, Zhang H, Jiang M, et al. Radiomics in ophthalmology: a systematic review. European Radiology. 2025;35(1):542–557. doi:10.1007/s00330-024-10911-4.",
"Carrera-Escalé L, Benali A, Rathert AC, et al. Radiomics-Based Assessment of OCT Angiography Images for Diabetic Retinopathy Diagnosis. Ophthalmology Science. 2023;3(2):100259. doi:10.1016/j.xops.2022.100259.",
"Perkins SW, Shah N, Whitney J, et al. Radiomic Characterization and Automated Classification of Drusen Substructure Phenotype Associated with High-Risk Dry Age-Related Macular Degeneration. Diagnostics. 2025;15(20):2594. doi:10.3390/diagnostics15202594.",
"Su Y, Xu X, Zuo P, et al. Value of MR-based radiomics in differentiating uveal melanoma from other intraocular masses in adults. European Journal of Radiology. 2020;131:109268.",
"Zhang H, Jiang M, Chan HC, et al. Whole-orbit radiomics: machine learning-based multi- and fused-region radiomics signatures for intravenous glucocorticoid response prediction in thyroid eye disease. Journal of Translational Medicine. 2024;22:56. doi:10.1186/s12967-023-04792-2.",
"Yedekci Y, Arimura H, Jin Y, et al. Non-invasive prediction of the secondary enucleation risk in uveal melanoma based on pretreatment CT and MRI prior to stereotactic radiotherapy. Strahlentherapie und Onkologie. 2026;202(5):476–484. doi:10.1007/s00066-025-02449-1.",
"Wang S, He X, Jian Z, et al. Advances and prospects of multi-modal ophthalmic artificial intelligence based on deep learning: a review. Eye and Vision. 2024;11:38. doi:10.1186/s40662-024-00405-1.",
"Zwanenburg A, Vallières M, Abdalah MA, et al. The Image Biomarker Standardization Initiative: Standardized Quantitative Radiomics for High-Throughput Image-based Phenotyping. Radiology. 2020;295(2):328–338. doi:10.1148/radiol.2020191145.",
"Tejani AS, Klontzas ME, Gatti AA, et al. Checklist for Artificial Intelligence in Medical Imaging (CLAIM): 2024 Update. Radiology: Artificial Intelligence. 2024;6(4):e240300.",
"Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378.",
"Moons KGM, et al. PROBAST+AI: an updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ. 2025;388:bmj-2024-082505."
]
for ref in refs: add_p(doc,ref,9.0,False,False,False,2)

# QA
doc.core_properties.title="影像组学与眼科疾病：高影响力综述框架与图表整合版"
doc.core_properties.author="ChatGPT"
doc.save(DOCX)
check=Document(DOCX)
assert len(check.tables)==4
assert len(check.inline_shapes)==7
assert len(check.paragraphs)>80
print(DOCX)
print("paragraphs",len(check.paragraphs),"tables",len(check.tables),"figures",len(check.inline_shapes),"bytes",DOCX.stat().st_size)
